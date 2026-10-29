"""Tes alur utama: upload master SAP, import Excel, input HP, ubah qty manual, opname, koreksi, laporan."""
import datetime as dt
import io
import json
import re
import shutil
import uuid
from pathlib import Path

import openpyxl
import pytest
from PIL import Image

from app import create_app

ROOT = Path(__file__).resolve().parent.parent
HARI_INI = dt.date.today()


def xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def supplementary_xlsx(include_plasma=False):
    """Workbook kecil yang mengikuti susunan sheet Supplementary PM produksi."""
    wb = openpyxl.Workbook()
    block = wb.active
    block.title = "BLOCK"
    block.append([None, None, "Company Code", "Group Estate", "Estate", "Division", "Block", "Block Name", "Phase"])
    block.append([None, None, "CM", "KALBAR", "CM3A", "1", "B-034", "D24", "TM"])
    block.append([None, None, "CM", "KALBAR", "CM3A", "2", "B-091", "E31", "TM"])
    block.append([None, None, "CM", "KALBAR", "CM3A", "3", "B-092", "G17R", "TM"])
    if include_plasma:
        block.append([None, None, "CM", "KALBAR", "CM9A", "3", "B-042", "E05P", "TM"])
        block.append([None, None, "CM", "KALBAR", "CM9A", "4", "B-138", "J16P", "TM"])
        block.append([None, None, "CM", "KALBAR", "CM9A", "4", "B-140", "J17P", "TM"])
    vra = wb.create_sheet("VRA ORDER")
    vra.append(["Order", "Equipment", "Description", "Main WorkCtr", "Planning Plant", "Description", "Order Type"])
    vra.append(["40001234", "EQ-01", "DT32 - DUMP TRUCK", "GOH-VRA", "CM3A", "VRA - OWN - EQ-01", "ZVRA"])
    vra.append(["40005678", "EQ-02", "EXA17 - EXCAVATOR", "GOH-VRA", "CM3A", "VRA - OWN - EQ-02", "ZVRA"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def jpeg_file(nama="bukti.jpg", warna=(235, 240, 245)):
    buf = io.BytesIO()
    Image.new("RGB", (640, 360), warna).save(buf, "JPEG", quality=90)
    buf.seek(0)
    return buf, nama


def jpeg_besar(nama="foto-besar.jpg"):
    buf = io.BytesIO()
    Image.effect_noise((2400, 1800), 45).convert("RGB").save(buf, "JPEG", quality=95)
    ukuran_asli = buf.tell()
    buf.seek(0)
    return (buf, nama), ukuran_asli


def jpeg_dengan_exif(nama="foto-metadata.jpg", waktu="2026:10:15 07:08:09"):
    buf = io.BytesIO()
    exif = Image.Exif()
    exif[36867] = waktu  # DateTimeOriginal
    Image.new("RGB", (640, 360), (235, 240, 245)).save(buf, "JPEG", quality=90, exif=exif)
    buf.seek(0)
    return buf, nama


MASTER_HEAD = ["Name 1", "Material Group", "Material", "Material description", "Plant", "Storage location", "Batch",
               "Unrestricted", "Base Unit of Measure", "Value Unrestricted", "Special Stock", "Special stock number",
               "Stock in Transit", "Material type"]

MB51_HEAD = ["Material Document", "Posting Date", "WBS element", "Material", "Plant", "Storage location",
             "Movement type", "Quantity", "Purchase order", "Amt.in Loc.Cur.", "G/L Account", "Vendor",
             "Material description", "Special Stock", "Text", "Cost Center", "User Name"]


def master_rows(qty_racun=100):
    return [
        MASTER_HEAD,
        ["KEBUN CMA 1 INTI", "904", "904.001.002", "RACUN ROUNDUP 1L", "CM3A", "ST01", "STORE", qty_racun, "L",
         qty_racun * 50000, None, None, 0, "SPAR"],
        ["KEBUN CMA 1 INTI", "905", "905.001.001", "BERAS 25KG", "CM3A", "ST01", "STORE", 40, "SAK", 40 * 300000,
         None, None, 0, "SPAR"],
        # special stock Q untuk material yang sama -> baris terpisah
        ["KEBUN CMA 1 INTI", "904", "904.001.002", "RACUN ROUNDUP 1L", "CM3A", "ST01", "STORE", 10, "L", 500000,
         "Q", "WBS-001", 0, "SPAR"],
        # Material yang sama di ST01 dan STTF tetap menjadi satu saldo stok gudang.
        ["KEBUN CMA 1 INTI", "906", "906.002.003", "EGREK", "CM3A", "ST01", "STORE", 5, "PC", 500000, None, None,
         0, "SPAR"],
        ["KEBUN CMA 1 INTI", "906", "906.002.003", "EGREK", "CM3A", "STTF", "STORE", 3, "PC", 300000, None, None,
         0, "SPAR"],
        [None, None, None, None, None, None, None, 158, None, 99999999, None, None, None, None],  # baris total SAP
    ]


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("STOK_DATA_DIR", str(tmp_path))
    app = create_app({"TESTING": True, "SECRET_KEY": "x"})
    c = app.test_client()
    c.app = app
    r = c.post("/login", data={"username": "admin", "password": "admin123"})
    assert r.status_code == 302
    c.get("/")  # set csrf
    with c.session_transaction() as s:
        c.csrf = s["csrf"]
    return c


def post(c, url, data=None, **kw):
    data = dict(data or {})
    if url == "/input" and data.get("tombol", "keluar") == "keluar":
        data.setdefault("tujuan_tipe", "blok")
        data.setdefault("tujuan_kode", "D24")
    data["_csrf"] = c.csrf
    return c.post(url, data=data, **kw)


def q(c, sql, args=()):
    import sqlite3
    db = sqlite3.connect(c.app.config["DATABASE"])
    db.row_factory = sqlite3.Row
    try:
        return db.execute(sql, args).fetchall()
    finally:
        db.close()


def stok(c, material, sp=""):
    return q(c, "SELECT stok FROM v_stok WHERE material = ? AND special_stock = ?", (material, sp))[0]["stok"]


def upload_master(c, rows, **extra):
    return post(c, "/master/upload", {"file": (xlsx(rows), "sap.xlsx"), "tgl_stok_awal": "2026-10-13", **extra},
                content_type="multipart/form-data", follow_redirects=True)


def mobile_login(c, device_uuid=None):
    """Login API APK dan kembalikan header Bearer untuk pengujian sinkron."""
    device_uuid = device_uuid or str(uuid.uuid4())
    response = c.post("/api/mobile/v1/login", json={
        "username": "admin",
        "password": "admin123",
        "device_uuid": device_uuid,
        "device_name": "HP Uji Gudang",
        "app_version": "0.1.0",
    })
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.get_json()['token']}", "X-App-Version": "0.1.0"}


def test_csrf_wajib(client):
    r = client.post("/master/upload", data={})
    assert r.status_code == 400


def test_upload_master_template_asli(client):
    with open(ROOT / "TEMPLATE_MASTER_SAP.xlsx", "rb") as f:
        r = post(client, "/master/upload", {"file": (f, "TEMPLATE_MASTER_SAP.xlsx")},
                 content_type="multipart/form-data", follow_redirects=True)
    assert r.status_code == 200
    rows = q(client, "SELECT * FROM materials")
    assert len(rows) == 1 and rows[0]["material"] == "904.001.002" and rows[0]["stok_awal"] == 10
    assert q(client, "SELECT nama FROM kebun WHERE kode='CM3A'")[0]["nama"] == "KEBUN CMA 1 INTI"


def test_master_menggabungkan_st01_sttf_dan_mengabaikan_sloc_lain(client):
    rows = master_rows()
    rows.insert(-1, ["KEBUN CMA 1 INTI", "906", "906.002.003", "EGREK", "CM3A", "TP07", "TRANSIT",
                     99, "PC", 9_900_000, None, None, 0, "SPAR"])
    r = upload_master(client, rows)
    html = r.data.decode()
    egrek = q(client, "SELECT * FROM materials WHERE material='906.002.003' AND special_stock='' ")[0]
    assert egrek["stok_awal"] == 8 and egrek["sap_unrestricted"] == 8
    assert egrek["storage_location"] == "ST01+STTF"
    assert "1 baris di luar ST01/STTF tidak dihitung sebagai stok gudang (TP07: 1)" in html


def test_upload_master_sinkron_target_negatif_saat_stok_sudah_nol(client):
    """Target SAP negatif dibatasi ke nol tanpa membuat penyesuaian nol atau menggagalkan upload."""
    upload_master(client, master_rows(qty_racun=1))
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    r = post(client, "/input", {
        "tombol": "keluar", "material_id": racun, "qty": "1", "penerima": "Budi",
        "foto": jpeg_file(),
    }, content_type="multipart/form-data")
    assert r.status_code == 302 and stok(client, "904.001.002") == 0

    r = upload_master(client, master_rows(qty_racun=0), sinkron="1")

    assert r.status_code == 200
    assert b"Master SAP diproses" in r.data
    assert stok(client, "904.001.002") == 0
    assert not q(client, "SELECT id FROM transactions WHERE sumber = 'sinkron_sap'")


def test_input_mendeteksi_bon_identik_dan_memerlukan_konfirmasi(client):
    upload_master(client, master_rows(qty_racun=20))
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    dasar = {
        "tombol": "keluar", "material_id": [str(racun)], "qty": ["2"], "penerima": "Bobi",
        "keterangan": "Afd 3", "tanggal": HARI_INI.isoformat(), "tujuan_tipe": "blok",
        "tujuan_kode": "D24",
    }
    r = post(client, "/input", {**dasar, "foto": jpeg_file("awal.jpg")},
             content_type="multipart/form-data")
    assert r.status_code == 302

    cek = post(client, "/api/cek-duplikat", dasar)
    data = cek.get_json()
    assert cek.status_code == 200 and len(data["duplikat"]) == 1
    assert data["duplikat"][0]["penerima"] == "Bobi"
    assert data["duplikat"][0]["items"][0]["qty"] == "2"

    beda_tanggal = {**dasar, "tanggal": (HARI_INI - dt.timedelta(days=1)).isoformat()}
    cek_beda_tanggal = post(client, "/api/cek-duplikat", beda_tanggal)
    assert cek_beda_tanggal.status_code == 200
    assert cek_beda_tanggal.get_json()["duplikat"] == []

    r = post(client, "/input", {**dasar, "foto": jpeg_file("ulang.jpg")},
             content_type="multipart/form-data", follow_redirects=True)
    assert "Kemungkinan double input ditemukan" in r.data.decode()
    assert len(q(client, "SELECT id FROM transactions WHERE sumber='hp' AND batal=0")) == 1

    r = post(client, "/input", {**dasar, "duplikat_token": data["token"],
                                  "foto": jpeg_file("benar-berbeda.jpg")},
             content_type="multipart/form-data")
    assert r.status_code == 302
    assert len(q(client, "SELECT id FROM transactions WHERE sumber='hp' AND batal=0")) == 2


def test_supplementary_tujuan_dapat_dicari_dan_goh_staff_tersedia(client):
    r = post(client, "/master/supplementary",
             {"file": (supplementary_xlsx(), "supplementary.xlsx")},
             content_type="multipart/form-data", follow_redirects=True)
    assert r.status_code == 200 and "3 blok dan 2 unit" in r.data.decode()
    targets = q(client, "SELECT target_type, code, sap_code, plant FROM supplementary_targets ORDER BY code")
    assert [(t["target_type"], t["code"], t["sap_code"]) for t in targets] == [
        ("blok", "D24", "B-034"), ("unit", "DT32", "40001234"),
        ("blok", "E31", "B-091"), ("unit", "EXA17", "40005678"), ("blok", "G17R", "B-092")]

    hasil_blok = client.get("/api/tujuan?q=D24&tipe=blok&plant=CM3A").get_json()
    assert hasil_blok[0]["kode"] == "D24" and hasil_blok[0]["sap_kode"] == "B-034"
    hasil_staff = client.get("/api/tujuan?q=staff&plant=CM3A").get_json()
    assert hasil_staff == [{"detail": "Cost center CM3AZGO991", "exact_match": True, "kode": "CM3AZGO991",
                            "nama": "GOH-STAFF", "plant": "CM3A", "sap_kode": "CM3AZGO991",
                            "tipe": "staff"}]
    assert client.get("/api/tujuan?q=kantor&plant=CM3A").get_json()[0]["tipe"] == "office"
    hasil_g17r = client.get("/api/tujuan?q=blok%20G17R&plant=CM3A").get_json()
    assert hasil_g17r[0]["kode"] == "G17R" and hasil_g17r[0]["exact_match"] is True
    hasil_afd3 = client.get("/api/tujuan?q=AFD%203&plant=CM3A").get_json()
    assert [target["kode"] for target in hasil_afd3] == ["G17R"]
    assert hasil_afd3[0]["exact_match"] is False and "pilih blok" in hasil_afd3[0]["detail"]
    assert client.get("/api/tujuan?q=AFD&plant=CM3A").get_json() == []
    hasil_dt32 = client.get("/api/tujuan?q=DT32&plant=CM3A").get_json()
    assert hasil_dt32[0]["kode"] == "DT32" and hasil_dt32[0]["exact_match"] is True
    hasil_ex17 = client.get("/api/tujuan?q=EX17&plant=CM3A").get_json()
    assert hasil_ex17[0]["kode"] == "EXA17" and hasil_ex17[0]["exact_match"] is True


def test_satu_material_dapat_dibagi_ke_beberapa_blok_dalam_satu_bon(client):
    upload_master(client, master_rows(qty_racun=100))
    post(client, "/master/supplementary", {"file": (supplementary_xlsx(), "supplementary.xlsx")},
         content_type="multipart/form-data")
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    data = {
        "tombol": "keluar", "material_id": [str(racun), str(racun)], "qty": ["1", "2"],
        "penerima": "Bobi", "tanggal": HARI_INI.isoformat(),
        "tujuan_tipe": "blok", "tujuan_kode": "D24",
        "item_tujuan_tipe": ["blok", "blok"], "item_tujuan_kode": ["D24", "E31"],
        "item_aktivitas_kode": ["", ""], "foto": jpeg_file("bagi-blok.jpg"),
    }
    r = post(client, "/input", data, content_type="multipart/form-data")
    assert r.status_code == 302
    trx = q(client, "SELECT bon, qty, tujuan_tipe, tujuan_kode FROM transactions ORDER BY id")
    assert len(trx) == 2 and len({t["bon"] for t in trx}) == 1
    assert [(t["qty"], t["tujuan_tipe"], t["tujuan_kode"]) for t in trx] == [
        (1, "blok", "D24"), (2, "blok", "E31")]
    assert stok(client, "904.001.002") == 97

    # Pengaman stok menghitung total dua pecahan, bukan meloloskan tiap baris sendiri-sendiri.
    gagal = {**data, "qty": ["60", "50"], "foto": jpeg_file("terlalu-banyak.jpg")}
    r = post(client, "/input", gagal, content_type="multipart/form-data", follow_redirects=True)
    assert "total pembagian 110 melebihi stok tersedia 97" in r.data.decode()
    assert len(q(client, "SELECT id FROM transactions")) == 2

    di_luar_master = {**data, "material_id": [str(racun)], "qty": ["1"],
                       "item_tujuan_tipe": [""], "item_tujuan_kode": [""],
                       "tujuan_kode": "ZZ99", "foto": jpeg_file("tujuan-invalid.jpg")}
    r = post(client, "/input", di_luar_master, content_type="multipart/form-data", follow_redirects=True)
    assert "ZZ99&#39; tidak terdaftar di Master Supplementary" in r.data.decode()
    assert len(q(client, "SELECT id FROM transactions")) == 2


def test_tujuan_goh_staff_disimpan_dengan_cost_center_otomatis(client):
    upload_master(client, master_rows(qty_racun=10))
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    r = post(client, "/input", {
        "tombol": "keluar", "material_id": str(racun), "qty": "1", "penerima": "Staff Kantor",
        "tujuan_tipe": "staff", "foto": jpeg_file("staff.jpg"),
    }, content_type="multipart/form-data")
    assert r.status_code == 302
    trx = q(client, "SELECT tujuan_tipe, tujuan_kode FROM transactions")[0]
    assert dict(trx) == {"tujuan_tipe": "staff", "tujuan_kode": "CM3AZGO991"}


def test_fame_wajib_unit_atau_workshop_dan_damkar_wajib_aktivitas(client):
    rows = master_rows(qty_racun=20)
    rows[1][2] = "929.001.Z18"
    rows[1][3] = "FAME (FATTY ACID METHYL ESTERS)"
    rows[1][8] = "L"
    upload_master(client, rows)
    fame = q(client, "SELECT id FROM materials WHERE material='929.001.Z18' AND special_stock='' ")[0]["id"]
    dasar = {"tombol": "keluar", "material_id": fame, "qty": "2", "penerima": "Darul",
             "tanggal": HARI_INI.isoformat()}

    r = post(client, "/input", {**dasar, "tujuan_tipe": "blok", "tujuan_kode": "D24",
                                 "foto": jpeg_file("fame-blok.jpg")},
             content_type="multipart/form-data", follow_redirects=True)
    assert "FAME, BBM, dan sparepart wajib diarahkan" in r.data.decode()

    r = post(client, "/input", {**dasar, "tujuan_tipe": "workshop", "foto": jpeg_file("fame-ws.jpg")},
             content_type="multipart/form-data")
    assert r.status_code == 302
    trx = q(client, "SELECT tujuan_tipe, tujuan_kode FROM transactions WHERE sumber='hp'")[0]
    assert trx["tujuan_tipe"] == "workshop" and trx["tujuan_kode"] == "CM3AZGO996"

    racun = q(client, "SELECT id FROM materials WHERE material='905.001.001' AND special_stock='' ")[0]["id"]
    damkar = {"tombol": "keluar", "material_id": racun, "qty": "1", "penerima": "Regu Damkar",
              "tanggal": HARI_INI.isoformat(), "tujuan_tipe": "damkar", "tujuan_kode": "D24"}
    r = post(client, "/input", {**damkar, "foto": jpeg_file("damkar-gagal.jpg")},
             content_type="multipart/form-data", follow_redirects=True)
    assert "Pilih aktivitas Damkar" in r.data.decode()
    r = post(client, "/input", {**damkar, "aktivitas_kode": "331", "foto": jpeg_file("damkar.jpg")},
             content_type="multipart/form-data")
    assert r.status_code == 302
    trx = q(client, "SELECT tujuan_tipe, tujuan_kode, aktivitas_kode FROM transactions "
                    "WHERE penerima='Regu Damkar'")[0]
    assert dict(trx) == {"tujuan_tipe": "damkar", "tujuan_kode": "D24", "aktivitas_kode": "331"}


def test_rekonsiliasi_mb51_mengisi_dokumen_transaksi_penyesuaian_dan_opname(client):
    upload_master(client, master_rows(qty_racun=10))
    ids = {r["material"]: r["id"] for r in q(client, "SELECT id, material FROM materials WHERE special_stock = ''")}
    racun, beras = ids["904.001.002"], ids["905.001.001"]
    foto = jpeg_file()
    post(client, "/input", {"tombol": "keluar", "material_id": [racun], "qty": ["3"],
                             "penerima": "Budi", "keterangan": "Afd 3", "foto": foto},
         content_type="multipart/form-data")
    post(client, f"/stok/{racun}/sesuaikan", {"qty_baru": "6", "alasan": "Koreksi salah input",
                                               "keterangan": "koreksi gudang"})
    post(client, "/opname", {f"fisik_{beras}": "38", "jenis_barang": "SPAREPART",
                              "alasan": "Stock opname (hitung fisik)"})
    assert stok(client, "904.001.002") == 6 and stok(client, "905.001.001") == 38

    rows = [
        MB51_HEAD,
        ["4900001001", HARI_INI, "", "904.LAMA.002", "CM3A", "ST01", "221", -3, "", 0, "", "",
         "RACUN ROUNDUP 1L", "", "Afd 3", "", "SAPUSER"],
        ["4900001002", HARI_INI, "", "904.001.002", "CM3A", "ST01", "702", -1, "", 0, "", "",
         "RACUN ROUNDUP 1L", "", "koreksi gudang", "", "SAPUSER"],
        ["4900001003", HARI_INI, "", "905.001.001", "CM3A", "ST01", "702", -2, "", 0, "", "",
         "BERAS 25KG", "", "stock opname", "", "SAPUSER"],
        [None, HARI_INI, None, None, None, None, None, -6, None, None, None, None, None, None, None, None, None],
    ]
    r = post(client, "/rekonsiliasi-sap", {"file": (xlsx(rows), "mb51.xlsx")},
             content_type="multipart/form-data")
    html = r.data.decode()
    assert r.status_code == 200 and "Siap diberi nomor SAP" in html
    assert "Kode berbeda" in html and "Stock Opname #" in html
    token = re.search(r'name="token" value="([0-9a-f]+)"', html).group(1)

    r = post(client, "/rekonsiliasi-sap/konfirmasi", {"token": token, "matches": ["0", "1", "2"]},
             follow_redirects=True)
    assert "3 transaksi diberi nomor dokumen SAP" in r.data.decode()
    trx = q(client, "SELECT jenis, sap_posted, sap_doc, sap_movement_type, sap_posting_date, sap_posted_at FROM transactions "
                    "WHERE sap_doc IS NOT NULL ORDER BY sap_doc")
    assert [t["sap_doc"] for t in trx] == ["4900001001", "4900001002", "4900001003"]
    assert all(t["sap_posted"] and t["sap_posting_date"] == HARI_INI.isoformat() for t in trx)
    assert all(t["sap_posted_at"].endswith("23:59:59") for t in trx)
    assert [t["sap_movement_type"] for t in trx] == ["221", "702", "702"]
    assert [t["jenis"] for t in trx] == ["keluar", "penyesuaian", "penyesuaian"]
    assert stok(client, "904.001.002") == 6 and stok(client, "905.001.001") == 38


def test_rekonsiliasi_mb51_wbs_biaya_bukan_special_stock(client):
    """WBS pada issue 221 normal tidak boleh memisahkan material unrestricted."""
    upload_master(client, master_rows(qty_racun=10))
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    post(client, "/input", {
        "tombol": "keluar", "material_id": racun, "qty": "3", "penerima": "Bobi",
        "keterangan": "Afd 3", "foto": jpeg_file("bobi.jpg"),
    }, content_type="multipart/form-data")

    rows51 = [MB51_HEAD,
              ["4900002210", HARI_INI, "X-CM3A1034.2.271", "904.001.002", "CM3A", "ST01",
               "221", -3, "", 0, "62100153", "", "RACUN ROUNDUP 1L", "", "Afd 3", "", "SAPUSER"]]
    r = post(client, "/rekonsiliasi-sap", {"mb51": (xlsx(rows51), "mb51-wbs.xlsx")},
             content_type="multipart/form-data")
    html = r.data.decode()
    assert r.status_code == 200 and "cocok otomatis" in html
    assert "Qty aplikasi tidak sama" not in html


def test_rekonsiliasi_memakai_kode_unit_untuk_membedakan_fame_qty_sama():
    """Teks unit SAP memilih transaksi FAME yang benar saat tanggal dan qty identik."""
    from app.views_admin import _pilih_kandidat

    candidates = [
        {"id": 11, "tujuan_kode": "DT07", "keterangan": "", "penerima": "Budi", "nomor_dokumen": ""},
        {"id": 12, "tujuan_kode": "DT32", "keterangan": "", "penerima": "Budi", "nomor_dokumen": ""},
    ]
    assert _pilih_kandidat(candidates, {"teks_sap": "dt32"})["id"] == 12
    assert _pilih_kandidat(candidates, {"teks_sap": "DT07 KM 12345"})["id"] == 11


def test_rekonsiliasi_mengabaikan_riwayat_sebelum_aplikasi(client):
    """Baris sebelum tanggal mulai diserap saldo awal, bukan dibuat sebagai transaksi aplikasi."""
    upload_master(client, master_rows(qty_racun=10))
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    post(client, "/input", {
        "tombol": "keluar", "material_id": racun, "qty": "1", "penerima": "Bobi",
        "keterangan": "Afd 3", "foto": jpeg_file("bobi.jpg"),
    }, content_type="multipart/form-data")
    kemarin = HARI_INI - dt.timedelta(days=1)
    rows51 = [MB51_HEAD,
              ["4900000014", kemarin, "", "904.001.002", "CM3A", "ST01", "221", -4,
               "", 0, "", "", "RACUN ROUNDUP 1L", "", "sebelum aplikasi", "", "SAPUSER"],
              ["4900000015", HARI_INI, "", "904.001.002", "CM3A", "ST01", "221", -1,
               "", 0, "", "", "RACUN ROUNDUP 1L", "", "Afd 3", "", "SAPUSER"]]
    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx(rows51), "mb51.xlsx"), "mb52": (xlsx(master_rows(qty_racun=5)), "mb52.xlsx"),
        "sap_per": f"{HARI_INI.isoformat()}T16:00", "same_day_included": "1",
        "mulai_rekonsiliasi": HARI_INI.isoformat(),
    }, content_type="multipart/form-data")
    html = r.data.decode()
    assert r.status_code == 200 and "Sebelum aplikasi" in html
    assert "diabaikan sebagai saldo awal" in html and "4900000014" not in html
    assert "4900000015" in html and "cocok otomatis" in html


def test_rekonsiliasi_bisa_memindahkan_transaksi_ke_plant_mb52_baru(client):
    """Bobi dapat dipasangkan ke issue CM9A dan material plant baru dibuat saat dikonfirmasi."""
    upload_master(client, master_rows(qty_racun=10))
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND plant='CM3A' "
                      "AND special_stock='' ")[0]["id"]
    post(client, "/input", {
        "tombol": "keluar", "material_id": racun, "qty": "1", "penerima": "Bobi",
        "keterangan": "Afd 3", "foto": jpeg_file("bobi.jpg"),
    }, content_type="multipart/form-data")
    rows51 = [MB51_HEAD,
              ["4900002215", HARI_INI, "X-CM9A3069.2.271", "904.001.002", "CM9A", "ST01",
               "221", -1, "", 0, "", "", "RACUN ROUNDUP 1L", "", "Afd 3", "", "SAPUSER"]]
    rows52_cm9a = [MASTER_HEAD,
                   ["KEBUN CMA 1 PLASMA", "904", "904.001.002", "RACUN ROUNDUP 1L", "CM9A", "ST01",
                    "STORE", 0, "L", 0, None, None, 0, "SPAR"]]
    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx(rows51), "mb51.xlsx"),
        "mb52": [(xlsx(master_rows(qty_racun=9)), "mb52-cm3a.xlsx"),
                 (xlsx(rows52_cm9a), "mb52-cm9a.xlsx")],
        "sap_per": f"{HARI_INI.isoformat()}T16:00", "same_day_included": "1",
        "mulai_rekonsiliasi": HARI_INI.isoformat(),
    }, content_type="multipart/form-data")
    html = r.data.decode()
    assert r.status_code == 200 and "Plant berbeda" in html
    assert "CM3A" in html and "CM9A" in html and "4900002215" in html
    token = re.search(r'name="token" value="([0-9a-f]+)"', html).group(1)
    r = post(client, "/rekonsiliasi-sap/konfirmasi", {"token": token, "matches": ["0"]},
             follow_redirects=True)
    assert "1 transaksi diberi nomor dokumen SAP" in r.data.decode()
    trx = q(client, """SELECT t.*, m.plant FROM transactions t JOIN materials m ON m.id=t.material_id
                       WHERE t.sap_doc='4900002215'""")[0]
    assert trx["plant"] == "CM9A" and trx["qty"] == 1 and trx["sap_posted"] == 1
    assert q(client, "SELECT id FROM materials WHERE plant='CM9A' AND material='904.001.002'")


def test_rekonsiliasi_cm9a_tanpa_mb52_menautkan_transaksi_flow_through(client):
    """CM9A bersaldo nol boleh menautkan dokumen SAP tanpa memindahkan material CM3A."""
    upload_master(client, master_rows(qty_racun=10))
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND plant='CM3A' "
                      "AND special_stock='' ")[0]["id"]
    post(client, "/input", {
        "tombol": "keluar", "material_id": racun, "qty": "1", "penerima": "Bobi",
        "keterangan": "Afd 3", "foto": jpeg_file("bobi.jpg"),
    }, content_type="multipart/form-data")
    rows51 = [MB51_HEAD,
              ["4900002216", HARI_INI, "X-CM9A3069.2.271", "904.001.002", "CM9A", "ST01",
               "221", -1, "", 0, "", "", "RACUN ROUNDUP 1L", "", "Afd 3", "", "SAPUSER"]]
    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx(rows51), "mb51-cm9a.xlsx"),
        "mb52": (xlsx(master_rows(qty_racun=9)), "mb52-cm3a.xlsx"),
        "sap_per": f"{HARI_INI.isoformat()}T16:00", "same_day_included": "1",
    }, content_type="multipart/form-data")
    html = r.data.decode()
    assert r.status_code == 200
    assert "Plant plasma tanpa saldo mengendap" in html
    assert 'name="matches"' in html
    token = re.search(r'name="token" value="([0-9a-f]+)"', html).group(1)
    r = post(client, "/rekonsiliasi-sap/konfirmasi", {"token": token, "matches": ["0"]},
             follow_redirects=True)
    assert "1 transaksi diberi nomor dokumen SAP" in r.data.decode()
    trx = q(client, "SELECT t.sap_posted, t.sap_doc, m.plant FROM transactions t "
                    "JOIN materials m ON m.id=t.material_id WHERE t.batal=0")[0]
    assert trx["sap_posted"] == 1 and trx["sap_doc"] == "4900002216" and trx["plant"] == "CM3A"
    assert not q(client, "SELECT id FROM materials WHERE plant='CM9A' AND material='904.001.002'")


def test_rekonsiliasi_plant_biasa_tanpa_mb52_tujuan_tetap_diblokir(client):
    """Plant selain flow-through tetap membutuhkan MB52 sebelum transaksi dipindahkan."""
    upload_master(client, master_rows(qty_racun=10))
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND plant='CM3A' "
                      "AND special_stock='' ")[0]["id"]
    post(client, "/input", {
        "tombol": "keluar", "material_id": racun, "qty": "1", "penerima": "Bobi",
        "keterangan": "Afd 3", "foto": jpeg_file("bobi.jpg"),
    }, content_type="multipart/form-data")
    rows51 = [MB51_HEAD,
              ["4900002217", HARI_INI, "X-CM4A3069.2.271", "904.001.002", "CM4A", "ST01",
               "221", -1, "", 0, "", "", "RACUN ROUNDUP 1L", "", "Afd 3", "", "SAPUSER"]]
    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx(rows51), "mb51-cm4a.xlsx"),
        "mb52": (xlsx(master_rows(qty_racun=9)), "mb52-cm3a.xlsx"),
        "sap_per": f"{HARI_INI.isoformat()}T16:00", "same_day_included": "1",
    }, content_type="multipart/form-data")
    html = r.data.decode()
    assert r.status_code == 200
    assert "MB52 CM4A belum tersedia" in html
    assert 'name="matches"' not in html


def test_rekonsiliasi_menjelaskan_qty_aplikasi_yang_tidak_ada_di_sap(client):
    upload_master(client, master_rows(qty_racun=10))
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    post(client, "/input", {
        "tombol": "keluar", "material_id": racun, "qty": "2", "penerima": "Bobi",
        "keterangan": "Afd 3", "foto": jpeg_file("bobi.jpg"),
    }, content_type="multipart/form-data")

    rows51 = [MB51_HEAD,
              ["4900002211", HARI_INI, "X-CM3A1034.2.271", "904.001.002", "CM3A", "ST01",
               "221", -1, "", 0, "62100153", "", "RACUN ROUNDUP 1L", "", "Afd 1", "", "SAPUSER"]]
    r = post(client, "/rekonsiliasi-sap", {"mb51": (xlsx(rows51), "mb51-beda-qty.xlsx")},
             content_type="multipart/form-data")
    html = r.data.decode()
    assert r.status_code == 200
    assert "Qty aplikasi tidak sama dengan baris SAP pada material dan tanggal ini." in html
    assert "-1 (MvT 221, dok. 4900002211)" in html


def test_rekonsiliasi_sttf_diperlakukan_sebagai_stok_gudang(client):
    """STTF yang sudah digabung dengan ST01 tidak boleh dianggap lokasi transit."""
    upload_master(client, master_rows(qty_racun=10))
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    post(client, "/input", {
        "tombol": "keluar", "material_id": racun, "qty": "1", "penerima": "Bobi",
        "keterangan": "Afd 3", "foto": jpeg_file("bobi.jpg"),
    }, content_type="multipart/form-data")

    rows51 = [MB51_HEAD,
              ["4900002212", HARI_INI, "X-CM3A1034.2.271", "904.001.002", "CM3A", "STTF",
               "221", -1, "", 0, "62100153", "", "RACUN ROUNDUP 1L", "", "Afd 3", "", "SAPUSER"]]
    r = post(client, "/rekonsiliasi-sap", {"mb51": (xlsx(rows51), "mb51-sttf.xlsx")},
             content_type="multipart/form-data")
    html = r.data.decode()
    assert r.status_code == 200
    assert "cocok otomatis" in html
    assert "Keputusan operator" not in html

def test_rekonsiliasi_mb51_mb52_menambah_transaksi_hilang_tanpa_stok_ganda(client):
    """Transaksi SAP yang hilang menggantikan sebagian sinkron lama; saldo akhir tetap MB52."""
    upload_master(client, master_rows(qty_racun=100))
    rows52 = master_rows(qty_racun=90)
    upload_master(client, rows52, sinkron="1", sap_per="2026-10-16T16:13")
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    assert stok(client, "904.001.002") == 90
    assert q(client, "SELECT qty FROM transactions WHERE material_id=? AND sumber='sinkron_sap' AND batal=0", (racun,))[0]["qty"] == -10

    rows51 = [MB51_HEAD,
              ["4900002001", dt.date(2026, 10, 15), "", "904.001.002", "CM3A", "ST01", "221", -4,
               "", 0, "", "", "RACUN ROUNDUP 1L", "", "langsung ke kantor", "", "SAPUSER"]]
    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx(rows51), "mb51.xlsx"), "mb52": (xlsx(rows52), "mb52.xlsx"),
        "sap_per": "2026-10-16T16:13",
    }, content_type="multipart/form-data")
    html = r.data.decode()
    assert "calon transaksi baru" in html and "Residual sinkron -10" in html
    token = re.search(r'name="token" value="([0-9a-f]+)"', html).group(1)
    r = post(client, "/rekonsiliasi-sap/konfirmasi", {"token": token, "news": ["0"]},
             follow_redirects=True)
    assert "1 transaksi baru ditambahkan" in r.data.decode()
    baru = q(client, "SELECT * FROM transactions WHERE sumber='sap_mb51' AND batal=0")
    assert len(baru) == 1 and baru[0]["jenis"] == "keluar" and baru[0]["qty"] == 4
    assert baru[0]["sap_posted"] == 1 and baru[0]["sap_doc"] == "4900002001"
    assert stok(client, "904.001.002") == 90
    sink = q(client, "SELECT qty FROM transactions WHERE material_id=? AND sumber='sinkron_sap' AND batal=0", (racun,))
    assert len(sink) == 1 and sink[0]["qty"] == -6
    assert q(client, "SELECT selisih_nyata FROM v_stok WHERE id=?", (racun,))[0]["selisih_nyata"] == 0


def test_rekonsiliasi_transfer_641_wajib_verifikasi_tujuan(client):
    """Movement transfer tidak boleh otomatis menjadi keluar biasa."""
    upload_master(client, master_rows(qty_racun=100))
    rows52 = master_rows(qty_racun=70)
    rows51 = [MB51_HEAD,
              ["4900002641", dt.date(2026, 10, 17), "", "904.001.002", "CM3A", "ST01", "641", -30,
               "", 0, "", "", "RACUN ROUNDUP 1L", "", "CMA 2", "", "SAPUSER"]]
    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx(rows51), "mb51-transfer.xlsx"), "mb52": (xlsx(rows52), "mb52.xlsx"),
        "sap_per": "2026-10-17T16:00", "same_day_included": "1",
    }, content_type="multipart/form-data")
    html = r.data.decode()
    assert "Keputusan operator" in html and "Transfer keluar ke kebun lain" in html
    assert "ST01" in html and "MvT 641" in html
    token = re.search(r'name="token" value="([0-9a-f]+)"', html).group(1)

    r = post(client, "/rekonsiliasi-sap/konfirmasi", {"token": token, "news": ["0"]})
    assert r.status_code == 400 and "Verifikasi transfer belum lengkap" in r.data.decode()
    assert not q(client, "SELECT id FROM transactions WHERE sumber='sap_mb51'")

    r = post(client, "/rekonsiliasi-sap/konfirmasi", {
        "token": token, "news": ["0"], "arah_0": "transfer_keluar", "pihak_0": "CMA 2",
    }, follow_redirects=True)
    assert "1 pergerakan transfer diverifikasi" in r.data.decode()
    trx = q(client, "SELECT * FROM transactions WHERE sumber='sap_mb51' AND batal=0")
    assert len(trx) == 1 and trx[0]["jenis"] == "transfer_keluar" and trx[0]["qty"] == 30
    assert trx[0]["kebun_asal"] == "CM3A" and trx[0]["kebun_tujuan"] == "CMA 2"
    review = q(client, "SELECT * FROM sap_movement_reviews")
    assert len(review) == 1 and review[0]["decision"] == "transfer_keluar" and review[0]["pihak"] == "CMA 2"
    assert stok(client, "904.001.002") == 70


def test_rekonsiliasi_rantai_plasma_menjadi_satu_transaksi_blok(client):
    """641→101→221 CM9A menjadi satu transaksi keluar ke blok, bukan tiga mutasi stok."""
    upload_master(client, master_rows(qty_racun=100))
    post(client, "/master/supplementary", {
        "file": (supplementary_xlsx(include_plasma=True), "supplementary-plasma.xlsx"),
    }, content_type="multipart/form-data")

    rows51 = [
        MB51_HEAD,
        ["4900002643", HARI_INI, "", "904.001.002", "CM3A", "ST01", "641", -4,
         "6007001007", 0, "", "", "RACUN ROUNDUP 1L", "", "", "", "SAPUSER"],
        ["5190000103", HARI_INI, "", "904.001.002", "CM9A", "ST01", "101", 4,
         "6007001007", 0, "", "", "RACUN ROUNDUP 1L", "", "PEMUPUKAN E05P", "", "SAPUSER"],
        ["4900002218", HARI_INI, "X-CM9A3042.2.293", "904.001.002", "CM9A", "ST01", "221", -4,
         "", 0, "62100154", "", "RACUN ROUNDUP 1L", "", "untuk pemupukan afd 3", "CM9AXMA410", "SAPUSER"],
    ]
    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx(rows51), "mb51-rantai-plasma.xlsx"),
        "mb52": (xlsx(master_rows(qty_racun=96)), "mb52-cm3a.xlsx"),
        "sap_per": f"{HARI_INI.isoformat()}T16:00", "same_day_included": "1",
    }, content_type="multipart/form-data")
    html = r.data.decode()
    assert r.status_code == 200
    assert "CM3A → CM9A → Blok E05P" in html
    assert html.count('name="news"') == 1
    assert html.count("tercakup rantai plasma") == 2
    assert "Keputusan operator" not in html

    token = re.search(r'name="token" value="([0-9a-f]+)"', html).group(1)
    r = post(client, "/rekonsiliasi-sap/konfirmasi", {"token": token, "news": ["0"]},
             follow_redirects=True)
    assert "1 transaksi baru ditambahkan" in r.data.decode()
    trx = q(client, """SELECT t.*, m.plant FROM transactions t
                        JOIN materials m ON m.id=t.material_id
                        WHERE t.sumber='sap_mb51' AND t.batal=0""")
    assert len(trx) == 1
    assert trx[0]["jenis"] == "keluar" and trx[0]["qty"] == 4 and trx[0]["plant"] == "CM3A"
    assert trx[0]["tujuan_tipe"] == "blok" and trx[0]["tujuan_kode"] == "E05P"
    assert all(doc in trx[0]["sap_doc"] for doc in ("4900002643", "5190000103", "4900002218"))
    docs = q(client, "SELECT * FROM transaction_sap_docs WHERE transaction_id=?", (trx[0]["id"],))
    assert len(docs) == 3
    assert stok(client, "904.001.002") == 96


def test_rekonsiliasi_rantai_plasma_terlambat_membagi_transaksi_per_blok(client):
    """MvT 101 yang terlambat memecah satu transaksi lama ke blok sesuai qty tanpa mengubah total."""
    upload_master(client, master_rows(qty_racun=100))
    racun = q(client, "SELECT id FROM materials WHERE plant='CM3A' AND material='904.001.002' "
                      "AND special_stock='' ")[0]["id"]
    post(client, "/input", {
        "tombol": "keluar", "tanggal": "2026-10-24", "material_id": str(racun), "qty": "6",
        "penerima": "A. Roni", "keterangan": "Pemupukan afd 4",
        "tujuan_tipe": "blok", "tujuan_kode": "J19", "foto": jpeg_file("pupuk.jpg"),
    }, content_type="multipart/form-data")
    trx_awal = q(client, "SELECT * FROM transactions WHERE sumber='hp' AND batal=0")[0]
    post(client, "/master/supplementary", {
        "file": (supplementary_xlsx(include_plasma=True), "supplementary-plasma.xlsx"),
    }, content_type="multipart/form-data")

    rows51 = [
        MB51_HEAD,
        ["4902536654", dt.date(2026, 10, 24), "", "904.001.002", "CM3A", "ST01", "641", -2.127,
         "6007001005", 0, "", "", "RACUN ROUNDUP 1L", "", "", "", "SAPUSER"],
        ["4902536654", dt.date(2026, 10, 24), "", "904.001.002", "CM3A", "ST01", "641", -3.873,
         "6007001005", 0, "", "", "RACUN ROUNDUP 1L", "", "", "", "SAPUSER"],
        ["5190854253", dt.date(2026, 10, 27), "", "904.001.002", "CM9A", "ST01", "101", 2.127,
         "6007001005", 0, "", "", "RACUN ROUNDUP 1L", "", "J16P", "", "SAPUSER"],
        ["5190854253", dt.date(2026, 10, 27), "", "904.001.002", "CM9A", "ST01", "101", 3.873,
         "6007001005", 0, "", "", "RACUN ROUNDUP 1L", "", "J17P", "", "SAPUSER"],
        ["4902536688", dt.date(2026, 10, 24), "X-CM9A5139.2.293", "904.001.002", "CM9A", "ST01",
         "221", -2.127, "", 0, "62100154", "", "RACUN ROUNDUP 1L", "", "", "CM9AXMA410", "SAPUSER"],
        ["4902536688", dt.date(2026, 10, 24), "X-CM9A5140.2.293", "904.001.002", "CM9A", "ST01",
         "221", -3.873, "", 0, "62100154", "", "RACUN ROUNDUP 1L", "", "", "CM9AXMA410", "SAPUSER"],
    ]
    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx(rows51), "mb51-rantai-terlambat.xlsx"),
        "mb52": (xlsx(master_rows(qty_racun=94)), "mb52-cm3a.xlsx"),
        "mulai_rekonsiliasi": "2026-10-24", "sap_per": "2026-10-28T08:00",
        "same_day_included": "1",
    }, content_type="multipart/form-data")
    html = r.data.decode()
    assert r.status_code == 200
    assert "1 transaksi → 2 blok plasma" in html
    assert "J16P" in html and "J17P" in html
    assert "MvT 101 baru diposting 27/10/2026" in html
    checkbox = re.search(r'name="matches" value="(\d+)"(?![^>]*checked)', html)
    assert checkbox

    token = re.search(r'name="token" value="([0-9a-f]+)"', html).group(1)
    r = post(client, "/rekonsiliasi-sap/konfirmasi", {
        "token": token, "matches": [checkbox.group(1)],
    }, follow_redirects=True)
    assert "1 transaksi dibagi mengikuti blok plasma" in r.data.decode()
    bagian = q(client, """SELECT t.*, m.plant FROM transactions t
                           JOIN materials m ON m.id=t.material_id
                           WHERE t.bon=? AND t.batal=0 ORDER BY t.tujuan_kode""", (trx_awal["bon"],))
    assert [(row["tujuan_kode"], row["qty"]) for row in bagian] == [("J16P", 2.127), ("J17P", 3.873)]
    assert all(row["plant"] == "CM3A" and row["sap_posted"] == 1 for row in bagian)
    assert sum(row["qty"] for row in bagian) == 6
    assert all(len(q(client, "SELECT id FROM transaction_sap_docs WHERE transaction_id=?", (row["id"],))) == 3
               for row in bagian)
    assert stok(client, "904.001.002") == 94


def test_rekonsiliasi_pasangan_641_101_meminta_verifikasi_ulang_keputusan_lama(client):
    """Ketika receipt plant tujuan tersedia, keputusan 641 lama yang bukan transfer harus diperbaiki manual."""
    upload_master(client, master_rows(qty_racun=100))
    rows52_cm3a = master_rows(qty_racun=96)
    keluar = ["4900002642", HARI_INI, "", "904.001.002", "CM3A", "ST01", "641", -4,
              "6007000991", 0, "", "", "RACUN ROUNDUP 1L", "", "CMA 1 PLASMA", "", "SAPUSER"]
    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx([MB51_HEAD, keluar]), "mb51-awal.xlsx"),
        "mb52": (xlsx(rows52_cm3a), "mb52-cm3a.xlsx"),
        "sap_per": f"{HARI_INI.isoformat()}T16:00", "same_day_included": "1",
    }, content_type="multipart/form-data")
    token = re.search(r'name="token" value="([0-9a-f]+)"', r.data.decode()).group(1)
    post(client, "/rekonsiliasi-sap/konfirmasi", {
        "token": token, "news": ["0"], "arah_0": "keluar_store",
    }, follow_redirects=True)
    transaction_id = q(client, "SELECT transaction_id FROM sap_movement_reviews")[0]["transaction_id"]
    assert q(client, "SELECT decision FROM sap_movement_reviews")[0]["decision"] == "keluar_store"

    rows52_cm9a = [MASTER_HEAD,
                   ["KEBUN CMA 1 PLASMA", "904", "904.001.002", "RACUN ROUNDUP 1L", "CM9A", "ST01",
                    "STORE", 4, "L", 0, None, None, 0, "SPAR"]]
    masuk = ["5190000101", HARI_INI, "", "904.001.002", "CM9A", "ST01", "101", 4,
             "6007000991", 0, "", "", "RACUN ROUNDUP 1L", "", "DARI CM3A", "", "SAPUSER"]
    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx([MB51_HEAD, keluar, masuk]), "mb51-lengkap.xlsx"),
        "mb52": [(xlsx(rows52_cm3a), "mb52-cm3a.xlsx"),
                 (xlsx(rows52_cm9a), "mb52-cm9a.xlsx")],
        "sap_per": f"{HARI_INI.isoformat()}T16:00", "same_day_included": "1",
    }, content_type="multipart/form-data")
    html = r.data.decode()
    assert "Verifikasi ulang transfer" in html and "CM3A → CM9A" in html
    assert "tercakup pasangan transfer" in html
    assert "hanya dapat ditandai transit" not in html
    token = re.search(r'name="token" value="([0-9a-f]+)"', html).group(1)
    r = post(client, "/rekonsiliasi-sap/konfirmasi", {
        "token": token, "reviews": ["0"], "arah_0": "transfer_keluar", "pihak_0": "CM9A",
    }, follow_redirects=True)
    assert "1 pergerakan transfer diverifikasi" in r.data.decode()
    trx = q(client, "SELECT jenis, kebun_tujuan, sap_doc FROM transactions WHERE id=?", (transaction_id,))[0]
    assert trx["jenis"] == "transfer_keluar" and trx["kebun_tujuan"] == "CM9A"
    assert "4900002642" in trx["sap_doc"] and "5190000101" in trx["sap_doc"]
    docs = q(client, "SELECT sap_doc, movement_type FROM transaction_sap_docs WHERE transaction_id=?",
             (transaction_id,))
    assert {(row["sap_doc"], row["movement_type"]) for row in docs} == {
        ("4900002642", "641"), ("5190000101", "101"),
    }
    reviews = q(client, "SELECT decision, pihak FROM sap_movement_reviews ORDER BY movement_type")
    assert {(row["decision"], row["pihak"]) for row in reviews} == {
        ("transfer_keluar", "CM9A"), ("transfer_masuk", "CM3A"),
    }


def test_rekonsiliasi_menandai_pasangan_221_222_sebagai_reversal(client):
    """Issue dan reversal SAP yang neto nol tidak ditawarkan sebagai transaksi baru."""
    upload_master(client, master_rows(qty_racun=100))
    rows51 = [
        MB51_HEAD,
        ["4902525560", dt.date(2026, 10, 24), "", "904.001.002", "CM3A", "ST01", "221", -6400,
         "", 0, "", "", "RACUN ROUNDUP 1L", "", "", "", "SAPUSER"],
        ["4902525606", dt.date(2026, 10, 24), "", "904.001.002", "CM3A", "ST01", "222", 6400,
         "", 0, "", "", "RACUN ROUNDUP 1L", "", "", "", "SAPUSER"],
    ]
    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx(rows51), "mb51-reversal.xlsx"),
        "mb52": (xlsx(master_rows(qty_racun=100)), "mb52.xlsx"),
        "sap_per": "2026-10-28T08:00", "same_day_included": "1",
    }, content_type="multipart/form-data")
    html = r.data.decode()
    assert r.status_code == 200
    assert "Reversal SAP" in html and html.count("dibatalkan di SAP") == 2
    assert "tidak ditemukan" not in html and 'name="news"' not in html
    assert not q(client, "SELECT id FROM transactions WHERE sumber='sap_mb51'")


def test_rekonsiliasi_tp07_bisa_ditetapkan_masih_transit_tanpa_ubah_stok(client):
    """101 di TP07 belum masuk Store, sehingga keputusan transit hanya menjadi jejak audit."""
    upload_master(client, master_rows(qty_racun=100))
    lengkap52 = master_rows(qty_racun=100)
    rows52 = [lengkap52[0], *lengkap52[2:]]  # material normal belum muncul pada saldo ST01 MB52
    rows51 = [MB51_HEAD,
              ["5000000101", dt.date(2026, 10, 15), "", "904.001.002", "CM3A", "TP07", "101", 10,
               "45000001", 0, "", "", "RACUN ROUNDUP 1L", "", "DARI LO", "", "SAPUSER"]]
    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx(rows51), "mb51-tp07.xlsx"), "mb52": (xlsx(rows52), "mb52.xlsx"),
        "sap_per": "2026-10-17T16:00", "same_day_included": "1",
    }, content_type="multipart/form-data")
    html = r.data.decode()
    assert "TP07" in html and "Masih transit — belum masuk ST01" in html
    assert "Material belum ada pada saldo MB52 ST01" in html
    token = re.search(r'name="token" value="([0-9a-f]+)"', html).group(1)
    r = post(client, "/rekonsiliasi-sap/konfirmasi", {
        "token": token, "news": ["0"], "arah_0": "transit",
    }, follow_redirects=True)
    assert "1 pergerakan transfer diverifikasi" in r.data.decode()
    assert not q(client, "SELECT id FROM transactions WHERE sumber='sap_mb51'")
    review = q(client, "SELECT decision, storage_location, transaction_id FROM sap_movement_reviews")
    assert len(review) == 1 and review[0]["decision"] == "transit" and review[0]["storage_location"] == "TP07"
    assert review[0]["transaction_id"] is None and stok(client, "904.001.002") == 100

    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx(rows51), "mb51-tp07.xlsx"), "mb52": (xlsx(rows52), "mb52.xlsx"),
        "sap_per": "2026-10-17T16:00", "same_day_included": "1",
    }, content_type="multipart/form-data")
    assert "sudah diverifikasi" in r.data.decode() and "Masih transit; stok ST01 tidak berubah" in r.data.decode()


def test_rekonsiliasi_dua_dokumen_sap_mengoreksi_qty_pemisah_ribuan(client):
    """7.633 input operator dikenali sebagai 7.633 kg ketika MB51 pecah menjadi dua dokumen."""
    upload_master(client, master_rows(qty_racun=10000))
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    post(client, "/input", {"tombol": "keluar", "material_id": racun, "qty": "7.633", "penerima": "A Roni",
                             "keterangan": "Afd 3", "foto": jpeg_file("x.jpg")},
         content_type="multipart/form-data")
    rows52 = master_rows(qty_racun=2367)
    upload_master(client, rows52, sinkron="1", sap_per=f"{HARI_INI.isoformat()}T16:13")
    rows51 = [MB51_HEAD,
              ["4900003001", HARI_INI, "", "904.001.002", "CM3A", "ST01", "221", -3000,
               "", 0, "", "", "RACUN ROUNDUP 1L", "", "Afd 3", "", "SAPUSER"],
              ["4900003002", HARI_INI, "", "904.001.002", "CM3A", "ST01", "221", -4633,
               "", 0, "", "", "RACUN ROUNDUP 1L", "", "Afd 3", "", "SAPUSER"]]
    r = post(client, "/rekonsiliasi-sap", {
        "mb51": (xlsx(rows51), "mb51.xlsx"), "mb52": (xlsx(rows52), "mb52.xlsx"),
        "sap_per": f"{HARI_INI.isoformat()}T16:13", "same_day_included": "1",
    }, content_type="multipart/form-data")
    html = r.data.decode()
    assert "2 dokumen → 1 transaksi" in html and "akan dikoreksi menjadi 7.633" in html
    token = re.search(r'name="token" value="([0-9a-f]+)"', html).group(1)
    r = post(client, "/rekonsiliasi-sap/konfirmasi", {"token": token, "matches": ["0"]},
             follow_redirects=True)
    assert "1 quantity salah format dikoreksi" in r.data.decode()
    t = q(client, "SELECT * FROM transactions WHERE material_id=? AND sumber='hp' AND batal=0", (racun,))[0]
    assert t["qty"] == 7633 and t["sap_doc"] == "4900003001, 4900003002" and t["sap_posted"] == 1
    assert len(q(client, "SELECT * FROM transaction_sap_docs WHERE transaction_id=?", (t["id"],))) == 2
    assert stok(client, "904.001.002") == 2367
    assert not q(client, "SELECT id FROM transactions WHERE material_id=? AND sumber='sinkron_sap' AND batal=0", (racun,))

def test_alur_lengkap(client):
    r = upload_master(client, master_rows())
    assert b"4 material baru" in r.data
    assert stok(client, "904.001.002") == 100
    assert stok(client, "904.001.002", "Q") == 10
    assert stok(client, "906.002.003") == 8  # 5 + 3 digabung
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock=''")[0]["id"]

    # kebun tujuan
    post(client, "/kebun", {"kode": "CM2B", "nama": "KEBUN CMA 2 PLASMA"})

    # --- input HP: KELUAR tanpa foto ditolak
    r = post(client, "/input", {"tombol": "keluar", "material_id": racun, "qty": "5", "penerima": "Budi"},
             content_type="multipart/form-data", follow_redirects=True)
    assert "Foto bukti wajib".encode() in r.data
    assert stok(client, "904.001.002") == 100

    foto = lambda: jpeg_file()  # noqa: E731
    # --- KELUAR melebihi stok ditolak
    r = post(client, "/input", {"tombol": "keluar", "material_id": racun, "qty": "101", "penerima": "Budi",
                                "foto": foto()}, content_type="multipart/form-data", follow_redirects=True)
    assert b"melebihi stok" in r.data
    # --- KELUAR ok (desimal koma)
    r = post(client, "/input", {"tombol": "keluar", "material_id": racun, "qty": "5,5", "penerima": "Budi",
                                "foto": foto()}, content_type="multipart/form-data", follow_redirects=True)
    assert b"tersimpan" in r.data
    assert stok(client, "904.001.002") == 94.5

    # --- TRANSFER kirim: wajib kebun & no dokumen
    r = post(client, "/input", {"tombol": "transfer", "arah": "kirim", "material_id": racun, "qty": "4",
                                "penerima": "Andi", "foto": foto()},
             content_type="multipart/form-data", follow_redirects=True)
    assert b"Nomor dokumen transfer wajib" in r.data and b"Kebun tujuan wajib" in r.data
    r = post(client, "/input", {"tombol": "transfer", "arah": "kirim", "material_id": racun, "qty": "4",
                                "penerima": "Andi", "kebun": "CM2B", "nomor_dokumen": "TRF-9", "foto": foto()},
             content_type="multipart/form-data", follow_redirects=True)
    assert b"tersimpan" in r.data
    t = q(client, "SELECT * FROM transactions WHERE jenis='transfer_keluar'")[0]
    assert t["kebun_asal"] == "CM3A" and t["kebun_tujuan"] == "CM2B" and t["foto"]
    assert stok(client, "904.001.002") == 90.5

    # --- MASUK dari HP tanpa foto boleh
    r = post(client, "/input", {"tombol": "masuk", "material_id": racun, "qty": "10"},
             content_type="multipart/form-data", follow_redirects=True)
    assert b"tersimpan" in r.data
    assert stok(client, "904.001.002") == 100.5

    # --- Ubah qty manual (penyesuaian) via fetch
    r = post(client, f"/stok/{racun}/sesuaikan", {"qty_baru": "97", "alasan": "Stock opname (hitung fisik)"},
             headers={"X-Requested-With": "fetch"})
    assert r.json["ok"], r.json
    assert stok(client, "904.001.002") == 97
    p = q(client, "SELECT * FROM transactions WHERE jenis='penyesuaian'")[0]
    assert p["qty"] == -3.5 and p["alasan"].startswith("Stock opname")
    r = post(client, f"/stok/{racun}/sesuaikan", {"qty_baru": "-1", "alasan": "x"},
             headers={"X-Requested-With": "fetch"})
    assert r.status_code == 400
    r = post(client, f"/stok/{racun}/sesuaikan", {"qty_baru": "5", "alasan": ""},
             headers={"X-Requested-With": "fetch"})
    assert r.status_code == 400

    # --- Koreksi qty transaksi KELUAR 5,5 -> 2,5 (stok naik 3)
    kel = q(client, "SELECT id FROM transactions WHERE jenis='keluar'")[0]["id"]
    post(client, f"/transaksi/{kel}/koreksi", {"qty": "2.5", "alasan": "salah ketik"})
    assert stok(client, "904.001.002") == 100
    # koreksi yang bikin minus ditolak
    masuk = q(client, "SELECT id FROM transactions WHERE jenis='masuk'")[0]["id"]
    r = post(client, f"/transaksi/{masuk}/koreksi", {"qty": "0.001", "alasan": "x"}, follow_redirects=True)
    assert stok(client, "904.001.002") == pytest.approx(90.001)
    # batal transaksi
    post(client, f"/transaksi/{masuk}/batal", {"alasan": "dobel"})
    assert stok(client, "904.001.002") == 90

    # --- Stock opname massal
    beras = q(client, "SELECT id FROM materials WHERE material='905.001.001'")[0]["id"]
    r = post(client, "/opname", {f"fisik_{beras}": "38", f"fisik_{racun}": "90", f"ket_{beras}": "2 sak basah",
                                 "alasan": "Stock opname (hitung fisik)", "jenis_barang": "sparepart",
                                 "lokasi": "CMA1"},
             follow_redirects=True)
    assert b"2 material dihitung, 1 disesuaikan" in r.data
    assert stok(client, "905.001.001") == 38
    # Berita Acara: material yang sama dengan kartu stok tetap tercatat
    sesi = q(client, "SELECT * FROM opname_sesi")[0]
    assert sesi["jenis_barang"] == "SPAREPART"
    items = {i["material_id"]: i for i in q(client, "SELECT * FROM opname_item")}
    assert items[racun]["trx_id"] is None and items[racun]["fisik"] == 90
    assert items[beras]["stok_buku"] == 40 and items[beras]["fisik"] == 38 and items[beras]["trx_id"]
    for url in ["/opname/riwayat", f"/opname/{sesi['id']}", "/opname/pengaturan"]:
        assert client.get(url).status_code == 200, url
    r = client.get(f"/opname/export?sesi={sesi['id']}")
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["SPAREPART", "Selisih"]
    ws = wb["SPAREPART"]
    assert ws["B6"].value == "BERITA ACARA PERHITUNGAN FISIK PERSEDIAAN" and ws["D8"].value == "SPAREPART"
    assert ws["L5"].value == ": PT CITRA MAHKOTA" and ws["L7"].value.startswith(": ")
    baris = {ws[f"C{r}"].value: r for r in (12, 13)}
    rb = baris["905.001.001"]
    assert ws[f"G{rb}"].value == 40 and ws[f"H{rb}"].value == 38 and ws[f"K{rb}"].value == 38
    assert ws[f"L{rb}"].value == 2 and ws[f"M{rb}"].value == 0  # Lebih = kartu stok > fisik
    assert "2 sak basah" in ws[f"N{rb}"].value and "Belum disesuaikan di SAP" in ws[f"N{rb}"].value
    assert ws["G16"].value == "Dibuat Oleh," and ws["G21"].value == "YUSUF ARISTOKRAT"
    assert ws.page_setup.orientation == "landscape"
    sel = wb["Selisih"]
    assert sel["E5"].value == "905.001.001" and sel["F6"].value is None

    # --- Import Excel: preview + konfirmasi
    tgl = HARI_INI.strftime("%d/%m/%Y")
    trx_bytes = xlsx([
        ["Tanggal (dd/mm/yyyy)", "Jenis Transaksi", "Material", "Quantity", "Base Unit of Measure", "Nomor Dokumen",
         "Nama Penerima", "Kebun Asal", "Kebun Tujuan", "Keterangan"],
        [tgl, "MASUK", "905.001.001", 10, "SAK", "DO-1", None, None, None, "Pembelian"],
        [tgl, "TRANSFER_MASUK", "905.001.001", 2, "SAK", "TRF-IN-1", None, "CMA 2 PLASMA", None, None],
        [tgl, "TRANSFER_KELUAR", "905.001.001", 5, "SAK", "TRF-OUT-1", None, "CMA 1 INTI", "KEBUN LAIN", None],
        ["Jenis: MASUK = ...", None, None, None, None, None, None, None, None, None],
    ]).getvalue()
    trx = io.BytesIO(trx_bytes)
    r = post(client, "/import", {"plant": "CM3A", "file": (trx, "trx.xlsx")}, content_type="multipart/form-data")
    html = r.data.decode()
    assert "1 error" in html and "belum terdaftar" in html  # KEBUN LAIN belum ada
    trx = io.BytesIO(trx_bytes)
    r = post(client, "/import", {"plant": "CM3A", "auto_kebun": "1", "file": (trx, "trx.xlsx")},
             content_type="multipart/form-data")
    html = r.data.decode()
    assert "error</span>" not in html
    token = re.search(r'name="token" value="([0-9a-f]+)"', html).group(1)
    r = post(client, "/import/konfirmasi", {"token": token}, follow_redirects=True)
    assert b"3 transaksi" in r.data
    assert stok(client, "905.001.001") == 38 + 10 + 2 - 5
    assert q(client, "SELECT kode FROM kebun WHERE kode='LAIN'")
    # import ulang file yang sama -> ditolak dobel
    trx = io.BytesIO(trx_bytes)
    r = post(client, "/import", {"plant": "CM3A", "auto_kebun": "1", "file": (trx, "trx.xlsx")},
             content_type="multipart/form-data")
    assert b"sudah pernah diinput" in r.data
    # KELUAR lewat excel ditolak
    r = post(client, "/import", {"plant": "CM3A", "file": (xlsx([
        ["Tanggal", "Jenis Transaksi", "Material", "Quantity"], [tgl, "KELUAR", "905.001.001", 1]]), "k.xlsx")},
        content_type="multipart/form-data")
    assert b"KELUAR hanya lewat input HP" in r.data

    def v(material="904.001.002"):
        return q(client, "SELECT * FROM v_stok WHERE material = ? AND special_stock = ''", (material,))[0]

    # --- Upload SAP ulang SEBELUM transaksi diposting (kasus nyata): SAP masih 100, aplikasi 90.
    # Bedanya hanya transaksi belum posting -> bukan selisih nyata, dan "samakan" tidak boleh mengembalikan stok.
    upload_master(client, master_rows(qty_racun=100))
    assert stok(client, "904.001.002") == 90
    assert v()["pending_sap"] == pytest.approx(-10)  # keluar 2,5 + transfer 4 + opname -3,5 (masuk dibatalkan)
    assert v()["selisih_nyata"] == pytest.approx(0)
    html = client.get("/selisih").data.decode()
    # racun & beras: ada opname/Ubah Qty (selisih fisik) + transaksi harian, semuanya belum diinput SAP
    assert "1. Selisih nyata (0)" in html
    assert "2. Selisih fisik belum disesuaikan di SAP (2)" in html
    assert "3. Transaksi harian belum diinput SAP (2)" in html
    assert v()["selisih_fisik"] == pytest.approx(-3.5) and v()["pending_rutin"] == pytest.approx(-6.5)
    assert b"904.001.002" in client.get("/stok?status=selisih").data  # selisih fisik = perlu tindakan
    post(client, "/selisih", {"ids": [racun]})
    assert stok(client, "904.001.002") == 90  # tidak berubah

    # --- SAP ternyata 80 -> selisih nyata 20; koreksi = SAP + belum posting = 70
    satu_jam_lalu = (dt.datetime.now() - dt.timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
    upload_master(client, master_rows(qty_racun=80), sap_per=satu_jam_lalu)
    assert v()["selisih_nyata"] == pytest.approx(20)
    assert b"904.001.002" in client.get("/stok?status=selisih").data
    post(client, "/selisih", {"ids": [racun]})
    assert stok(client, "904.001.002") == 70
    sink = q(client, "SELECT * FROM transactions WHERE sumber = 'sinkron_sap'")
    assert len(sink) == 1 and sink[0]["sap_posted"] == 1
    assert v()["selisih_nyata"] == pytest.approx(0)

    # --- Tandai posting SAP lalu upload SAP terbaru (70) -> tidak ada yang tertunda
    ids = [r["id"] for r in q(client, "SELECT id FROM transactions WHERE material_id = ? AND batal = 0", (racun,))]
    r = post(client, "/transaksi/posting", {"ids": ids, "aksi": "posting", "sap_doc": "4900001"},
             follow_redirects=True)
    assert b"ditandai SUDAH posting SAP" in r.data
    assert client.get("/transaksi?status=belum_posting&q=904.001.002").data.count(b'name="ids"') == 0
    # Ditandai ✓ setelah file SAP (80) dibuat -> SAP perkiraan = 80 - 10 = 70, tidak lagi "belum diinput"
    assert v()["pending_sap"] == 0 and v()["n_tunggu_upload"] == 3
    assert v()["sap_perkiraan"] == pytest.approx(70) and v()["selisih_nyata"] == pytest.approx(0)
    assert "904.001.002" not in client.get("/selisih").data.decode()
    # Export SAP terbaru ternyata 75 (input di SAP kurang) -> selisih nyata -5
    upload_master(client, master_rows(qty_racun=75))
    assert v()["n_tunggu_upload"] == 0 and v()["selisih_nyata"] == pytest.approx(-5)
    upload_master(client, master_rows(qty_racun=70))
    assert v()["pending_sap"] == 0 and v()["selisih_nyata"] == pytest.approx(0)

    # --- Export Excel
    r = client.get("/transaksi/export?q=904.001.002")
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    head = [c.value for c in ws[1]]
    assert "Posting SAP" in head and ws.max_row == 1 + len(q(
        client, "SELECT id FROM transactions WHERE material_id = ?", (racun,)))
    assert ws.cell(2, head.index("No. Dok SAP") + 1).value in ("4900001", None)
    r = client.get("/stok/export")
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    assert ws.cell(1, 14).value == "Selisih Nyata" and ws.max_row == 1 + 4

    # Selisih fisik beras (opname 40 -> 38) ditandai sudah disesuaikan di SAP dari halaman Selisih SAP
    beras_v = q(client, "SELECT * FROM v_stok WHERE material = '905.001.001'")[0]
    assert beras_v["n_selisih_fisik"] == 1
    r = post(client, "/selisih", {"aksi": "posting_fisik", "ids": [beras], "sap_doc": "MI-77"}, follow_redirects=True)
    assert "1 penyesuaian ditandai" in r.data.decode()
    beras_v = q(client, "SELECT * FROM v_stok WHERE material = '905.001.001'")[0]
    assert beras_v["n_selisih_fisik"] == 0 and beras_v["n_pending_rutin"] > 0  # transaksi harian tidak ikut ditandai

    # --- Halaman & laporan
    for url in ["/", "/stok", f"/stok/{racun}", "/transaksi", "/opname", "/master", "/import", "/kebun",
                "/kategori", "/pengguna", "/audit", "/laporan", "/input", "/akun", "/akses",
                "/transaksi?status=tanpa_foto", "/stok?status=selisih", "/api/material?q=racun"]:
        assert client.get(url).status_code == 200, url
    r = client.get(f"/laporan?dari={HARI_INI.replace(day=1)}&sampai={HARI_INI}&format=xlsx")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Ringkasan", "Transfer Antar Kebun", "Stok Menipis", "Rincian Mutasi"]
    assert wb["Transfer Antar Kebun"].max_row == 4  # header + transfer HP + transfer import masuk & keluar
    r = client.get(f"/laporan?dari={HARI_INI.replace(day=1)}&sampai={HARI_INI}&format=pdf")
    assert r.status_code == 200 and r.data[:4] == b"%PDF"


def test_peran_petugas(client):
    upload_master(client, master_rows())
    post(client, "/pengguna", {"username": "petugas1", "nama": "Petugas Satu", "password": "rahasia1",
                               "role": "petugas", "kebun_kode": "CM3A"})
    c = client.app.test_client()
    c.post("/login", data={"username": "petugas1", "password": "rahasia1"})
    assert c.get("/").status_code == 302  # diarahkan ke input
    assert c.get("/input").status_code == 200
    assert c.get("/master").status_code == 403
    assert c.get("/laporan").status_code == 403
    c.get("/input")
    with c.session_transaction() as s:
        csrf = s["csrf"]
    r = c.post("/stok/1/sesuaikan", data={"_csrf": csrf, "qty_baru": "1", "alasan": "x"})
    assert r.status_code == 403


def test_material_baru(client):
    """Barang baru pertama kali dipesan: tambah manual -> MASUK -> export SAP berikutnya tersambung ke baris yang sama."""
    upload_master(client, master_rows())
    page = client.get("/master/baru?material=907.100.001").data.decode()
    assert 'value="907.100.001"' in page and "jangan" in page
    r = post(client, "/master/baru", {"plant": "CM3A", "material": "907.100.001", "deskripsi": "Pompa air",
                                      "satuan": "pc", "stok_minimum": "1"})
    assert r.status_code == 302 and "tombol=masuk" in r.headers["Location"]
    m = q(client, "SELECT * FROM v_stok WHERE material = '907.100.001'")[0]
    assert m["stok"] == 0 and m["sap_unrestricted"] is None and m["satuan"] == "PC" and m["deskripsi"] == "POMPA AIR"
    assert b"POMPA AIR" in client.get(f"/input?tombol=masuk&material_id={m['id']}").data
    # duplikat ditolak
    r = post(client, "/master/baru", {"plant": "CM3A", "material": "907.100.001", "deskripsi": "x", "satuan": "PC"},
             follow_redirects=True)
    assert b"sudah ada" in r.data
    post(client, "/input", {"tombol": "masuk", "material_id": m["id"], "qty": "4", "nomor_dokumen": "DO-9"},
         content_type="multipart/form-data")
    assert stok(client, "907.100.001") == 4
    tid = q(client, "SELECT id FROM transactions WHERE material_id = ?", (m["id"],))[0]["id"]
    post(client, "/transaksi/posting", {"ids": [tid], "aksi": "posting"})
    # export SAP berikutnya memuat material tsb (GR 4 pc)
    rows = master_rows()
    rows.insert(1, ["KEBUN CMA 1 INTI", "907", "907.100.001", "POMPA AIR SANYO", "CM3A", "ST01", "STORE", 4, "PC",
                    4_000_000, None, None, 0, "SPAR"])
    upload_master(client, rows)
    semua = q(client, "SELECT * FROM v_stok WHERE material = '907.100.001'")
    assert len(semua) == 1 and semua[0]["id"] == m["id"]
    assert semua[0]["stok"] == 4 and semua[0]["sap_unrestricted"] == 4 and semua[0]["selisih_nyata"] == 0
    assert semua[0]["deskripsi"] == "POMPA AIR SANYO"


def test_bon_multi_barang(client):
    """Satu bon berisi banyak barang: satu foto, satu penerima, semua atau tidak sama sekali."""
    upload_master(client, master_rows())
    ids = {r["material"] + r["special_stock"]: r["id"] for r in q(client, "SELECT id, material, special_stock FROM materials")}
    racun, beras, egrek = ids["904.001.002"], ids["905.001.001"], ids["906.002.003"]
    foto = lambda: jpeg_file("bon.jpg")  # noqa: E731

    # satu barang melebihi stok -> seluruh bon ditolak
    r = post(client, "/input", {"tombol": "keluar", "material_id": [racun, beras, egrek], "qty": ["5", "2", "99"],
                                "penerima": "Mandor Budi", "foto": foto()},
             content_type="multipart/form-data", follow_redirects=True)
    html = r.data.decode()
    assert "EGREK: Qty 99 melebihi stok" in html and html.count('class="item-row"') == 3 + 1  # + template baris
    assert not q(client, "SELECT id FROM transactions")
    # barang dobel ditolak
    r = post(client, "/input", {"tombol": "keluar", "material_id": [racun, racun], "qty": ["1", "1"],
                                "penerima": "Budi", "foto": foto()},
             content_type="multipart/form-data", follow_redirects=True)
    assert "tercantum dua kali" in r.data.decode()

    r = post(client, "/input", {"tombol": "keluar", "material_id": [racun, beras, egrek], "qty": ["5", "2,5", "3"],
                                "penerima": "Mandor Budi", "keterangan": "afd 3", "foto": foto(),
                                "foto_2": jpeg_file("bon2.jpg", (245, 235, 220))},
             content_type="multipart/form-data", follow_redirects=True)
    assert "tersimpan, 3 barang" in r.data.decode()
    trx = q(client, "SELECT * FROM transactions ORDER BY id")
    assert len(trx) == 3 and len({t["bon"] for t in trx}) == 1 and trx[0]["bon"].startswith("BON-")
    assert len({t["foto"] for t in trx}) == 1 and trx[0]["foto"]  # satu foto dipakai bersama
    assert len({t["foto_2"] for t in trx}) == 1 and trx[0]["foto_2"]  # foto kedua juga satu per bon
    for kolom in ("foto", "foto_2"):
        path_foto = Path(client.app.config["UPLOAD_FOLDER"]) / trx[0][kolom]
        assert path_foto.suffix == ".jpg" and path_foto.exists()
        with Image.open(path_foto) as gambar:
            # Sumber tes berwarna terang; kotak gelap di kanan bawah membuktikan watermark sudah digambar.
            area = gambar.crop((gambar.width // 2, gambar.height * 2 // 3, gambar.width, gambar.height))
            assert sum(1 for p in area.get_flattened_data() if max(p) < 100) > 500
    assert {t["penerima"] for t in trx} == {"Mandor Budi"}
    halaman_trx = client.get(f"/transaksi?bon={trx[0]['bon']}").data.decode()
    assert "foto 1" in halaman_trx and "foto 2" in halaman_trx
    assert stok(client, "904.001.002") == 95 and stok(client, "905.001.001") == 37.5 and stok(client, "906.002.003") == 5
    # daftar transaksi bisa difilter per bon, export memuat kolom Bon
    bon = trx[0]["bon"]
    assert client.get(f"/transaksi?bon={bon}").data.decode().count('name="ids"') == 3
    ws = openpyxl.load_workbook(io.BytesIO(client.get(f"/transaksi/export?bon={bon}").data)).active
    assert ws["B1"].value == "Bon" and ws["B2"].value == bon and ws.max_row == 4

    # --- Edit keterangan / penerima / nomor dokumen (qty tidak berubah)
    assert 'data-trx-dialog="dlgEdit"' in client.get(f"/transaksi?bon={bon}").data.decode()
    t1 = trx[0]["id"]
    r = post(client, f"/transaksi/{t1}/edit", {"keterangan": "Afd 3 blok B", "penerima": "Mandor Budi S",
                                              "nomor_dokumen": "", "semua_bon": "1"}, follow_redirects=True)
    assert f"bon {bon} diperbarui" in r.data.decode()
    ubah = q(client, "SELECT * FROM transactions WHERE bon = ?", (bon,))
    assert {u["keterangan"] for u in ubah} == {"Afd 3 blok B"} and {u["penerima"] for u in ubah} == {"Mandor Budi S"}
    assert stok(client, "904.001.002") == 95  # qty tidak tersentuh
    post(client, f"/transaksi/{t1}/edit", {"keterangan": "khusus baris ini", "penerima": "Mandor Budi S"})
    assert [u["keterangan"] for u in q(client, "SELECT keterangan FROM transactions WHERE bon = ? ORDER BY id", (bon,))] \
        == ["khusus baris ini", "Afd 3 blok B", "Afd 3 blok B"]
    r = post(client, f"/transaksi/{t1}/edit", {"keterangan": "x", "penerima": ""}, follow_redirects=True)
    assert "Nama penerima wajib" in r.data.decode()
    assert len(q(client, "SELECT * FROM audit_log WHERE aksi = 'edit_transaksi'")) == 4

    # bon kedua mendapat nomor berikutnya; MASUK tanpa foto boleh

    post(client, "/input", {"tombol": "masuk", "material_id": [racun, beras], "qty": ["10", "1"],
                            "nomor_dokumen": "DO-5"}, content_type="multipart/form-data")
    bon2 = q(client, "SELECT DISTINCT bon FROM transactions WHERE jenis = 'masuk'")[0]["bon"]
    assert bon2 != bon and bon2.endswith("-002")
    # --- Satukan bon: dua KELUAR yang terlanjur beda bon
    post(client, "/input", {"tombol": "keluar", "material_id": [racun], "qty": ["1"], "penerima": "Sukardi",
                            "foto": foto()}, content_type="multipart/form-data")
    post(client, "/input", {"tombol": "keluar", "material_id": [beras], "qty": ["1"], "penerima": "Sukardi",
                            "foto": foto()}, content_type="multipart/form-data")
    a, b = q(client, "SELECT id, bon FROM transactions WHERE penerima = 'Sukardi' ORDER BY id")
    assert a["bon"] != b["bon"]
    stok_sebelum = stok(client, "904.001.002")
    r = post(client, "/transaksi/gabung-bon", {"ids": [a["id"], b["id"]]}, follow_redirects=True)
    assert f"disatukan ke bon {a['bon']}" in r.data.decode()
    assert {x["bon"] for x in q(client, "SELECT bon FROM transactions WHERE penerima = 'Sukardi'")} == {a["bon"]}
    assert stok(client, "904.001.002") == stok_sebelum  # stok tidak berubah
    # beda jenis ditolak
    masuk_id = q(client, "SELECT id FROM transactions WHERE jenis = 'masuk' LIMIT 1")
    if masuk_id:
        r = post(client, "/transaksi/gabung-bon", {"ids": [a["id"], masuk_id[0]["id"]]}, follow_redirects=True)
        assert "jenis transaksi berbeda" in r.data.decode()
    r = post(client, "/transaksi/gabung-bon", {"ids": [a["id"]]}, follow_redirects=True)
    assert "minimal 2" in r.data.decode()

    # --- Bug nomor ganda: setelah bon digabung, bon baru TIDAK boleh memakai nomor yang sudah ada
    semua_bon = {x["bon"] for x in q(client, "SELECT DISTINCT bon FROM transactions WHERE bon IS NOT NULL")}
    post(client, "/input", {"tombol": "keluar", "material_id": [egrek], "qty": ["1"], "penerima": "Satria",
                            "foto": foto()}, content_type="multipart/form-data")
    bon_satria = q(client, "SELECT bon FROM transactions WHERE penerima = 'Satria'")[0]["bon"]
    assert bon_satria not in semua_bon
    assert bon_satria != b["bon"]  # nomor bon yang hilang karena digabung tidak dipakai ulang (kode lama memberi nomor ini)
    # gabung beda penerima ditolak
    r = post(client, "/transaksi/gabung-bon", {"ids": [a["id"], q(client, "SELECT id FROM transactions WHERE penerima = 'Satria'")[0]["id"]]},
             follow_redirects=True)
    assert "nama penerima berbeda" in r.data.decode()
    # --- Pisahkan: satu barang Sukardi keluar dari bonnya menjadi bon baru
    r = post(client, "/transaksi/pisah-bon", {"ids": [b["id"]]}, follow_redirects=True)
    assert "dipisahkan menjadi bon baru" in r.data.decode()
    bon_b = q(client, "SELECT bon FROM transactions WHERE id = ?", (b["id"],))[0]["bon"]
    assert bon_b != a["bon"] and bon_b != bon_satria
    assert q(client, "SELECT bon FROM transactions WHERE id = ?", (a["id"],))[0]["bon"] == a["bon"]
    assert stok(client, "904.001.002") == stok_sebelum
    assert q(client, "SELECT * FROM audit_log WHERE aksi = 'pisah_bon'")

    # GET dengan material_id langsung menyiapkan barisnya

    assert client.get(f"/input?tombol=masuk&material_id={egrek}").data.decode().count('class="item-row"') == 1 + 1


def test_foto_kedua_opsional_dan_material_expired_diblokir(client):
    upload_master(client, master_rows())
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    foto = lambda: jpeg_file("foto.jpg")  # noqa: E731

    # Flag manual langsung memblokir input HP dan terlihat di pencarian/stok.
    r = post(client, "/master/atur", {"ids": [racun], "expired": "1"},
             headers={"X-Requested-With": "fetch"})
    assert r.json["ok"]
    api = client.get("/api/material?q=904.001.002").json[0]
    assert api["expired_effective"] is True
    assert "EXPIRED" in client.get("/stok?status=expired").data.decode()
    r = post(client, "/input", {"tombol": "keluar", "material_id": racun, "qty": "1",
                                 "penerima": "Budi", "foto": foto()},
             content_type="multipart/form-data", follow_redirects=True)
    assert "barang EXPIRED" in r.data.decode() and not q(client, "SELECT id FROM transactions")

    # Import Excel juga ditolak untuk material expired.
    r = post(client, "/import", {"plant": "CM3A", "file": (xlsx([
        ["Tanggal", "Jenis Transaksi", "Material", "Quantity"],
        [HARI_INI, "MASUK", "904.001.002", 1],
    ]), "expired.xlsx")}, content_type="multipart/form-data")
    assert "sudah EXPIRED" in r.data.decode()

    # Tanggal expired bekerja otomatis. Tanggal masa depan belum memblokir.
    kemarin = (HARI_INI - dt.timedelta(days=1)).isoformat()
    besok = (HARI_INI + dt.timedelta(days=1)).isoformat()
    post(client, "/master/atur", {"ids": [racun], "expired": "0", "expired_at": kemarin},
         headers={"X-Requested-With": "fetch"})
    assert client.get("/api/material?q=904.001.002").json[0]["expired_effective"] is True
    post(client, "/master/atur", {"ids": [racun], "expired_at": besok},
         headers={"X-Requested-With": "fetch"})

    # Foto kedua boleh kosong saat input, lalu dapat dilampirkan belakangan.
    r = post(client, "/input", {"tombol": "keluar", "material_id": racun, "qty": "1",
                                 "penerima": "Budi", "foto": foto()},
             content_type="multipart/form-data")
    assert r.status_code == 302
    t = q(client, "SELECT * FROM transactions")[0]
    assert t["foto"] and t["foto_2"] is None
    r = post(client, f"/transaksi/{t['id']}/foto", {"slot": "2", "foto": foto()},
             content_type="multipart/form-data", follow_redirects=True)
    assert "Foto 2 transaksi" in r.data.decode()
    assert q(client, "SELECT foto_2 FROM transactions WHERE id=?", (t["id"],))[0]["foto_2"]


def test_kode_fame_lama_dikunci_dan_kode_z_tetap_bisa_dipakai(client):
    """Kode FAME non-Z tetap terlihat untuk audit, tetapi tidak boleh dipakai pada transaksi baru."""
    rows = master_rows()
    rows[-1][7] += 266_000
    rows.insert(-1, [
        "KEBUN CMA 1 INTI", "929", "929.001.007", "FAME (FATTY ACID METHYL ESTERS)",
        "CM3A", "ST01", "STORE", 16_000, "L", 256_000_000, None, None, 0, "SPAR",
    ])
    rows.insert(-1, [
        "KEBUN CMA 1 INTI", "929", "929.001.Z18", "FAME (FATTY ACID METHYL ESTERS)",
        "CM3A", "ST01", "STORE", 250_000, "L", 4_000_000_000, None, None, 0, "SPAR",
    ])
    upload_master(client, rows)
    lama = q(client, "SELECT id FROM materials WHERE material='929.001.007'")[0]["id"]
    pengganti = q(client, "SELECT id FROM materials WHERE material='929.001.Z18'")[0]["id"]
    alasan = ("Dikunci sementara. Gunakan 929.001.Z18 sampai stok kode Z18 habis; "
              "buka kembali 929.001.007 setelah itu.")

    r = post(client, "/master/atur", {
        "ids": [lama], "input_locked": "1", "input_lock_reason": alasan,
    }, headers={"X-Requested-With": "fetch"})
    assert r.json["ok"]
    api_lama = next(m for m in client.get("/api/material?q=929.001.007").json
                    if m["material"] == "929.001.007")
    assert api_lama["input_locked"] is True and api_lama["input_lock_reason"] == alasan
    assert "DIKUNCI" in client.get("/stok?status=locked").data.decode()

    r = post(client, "/input", {
        "tombol": "masuk", "material_id": lama, "qty": "1", "nomor_dokumen": "UJI-FAME-LAMA",
    }, follow_redirects=True)
    assert alasan in r.data.decode()
    assert not q(client, "SELECT id FROM transactions WHERE material_id = ?", (lama,))

    r = post(client, "/import", {"plant": "CM3A", "file": (xlsx([
        ["Tanggal", "Jenis Transaksi", "Material", "Quantity", "Nomor Dokumen"],
        [HARI_INI, "MASUK", "929.001.007", 1, "UJI-IMPORT-FAME"],
    ]), "fame-lama.xlsx")}, content_type="multipart/form-data")
    assert "dikunci dari input" in r.data.decode() and alasan in r.data.decode()

    # Kode Z adalah kode operasional yang diposting SAP dan tidak ikut terkunci.
    r = post(client, "/input", {
        "tombol": "masuk", "material_id": pengganti, "qty": "1", "nomor_dokumen": "UJI-FAME-Z",
    })
    assert r.status_code == 302
    assert len(q(client, "SELECT id FROM transactions WHERE material_id = ?", (pengganti,))) == 1


def test_input_menerima_foto_langsung_dari_kamera(client):
    """Field kamera terpisah tetap tersimpan sebagai foto transaksi dan bon."""
    upload_master(client, master_rows())
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    r = post(client, "/input", {
        "tombol": "keluar", "material_id": racun, "qty": "1", "penerima": "Budi",
        "foto_camera": jpeg_file("kamera-orang.jpg"),
        "foto_2_camera": jpeg_file("kamera-bon.jpg", (235, 225, 210)),
    }, content_type="multipart/form-data")
    assert r.status_code == 302
    transaksi = q(client, "SELECT foto, foto_2 FROM transactions")[0]
    assert transaksi["foto"] and transaksi["foto_2"]
    assert (Path(client.app.config["UPLOAD_FOLDER"]) / transaksi["foto"]).exists()
    assert (Path(client.app.config["UPLOAD_FOLDER"]) / transaksi["foto_2"]).exists()


def test_foto_besar_dikompres_setelah_diberi_watermark(client):
    upload_master(client, master_rows())
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    foto, ukuran_asli = jpeg_besar()
    r = post(client, "/input", {"tombol": "keluar", "material_id": racun, "qty": "1",
                                 "penerima": "Budi", "foto": foto},
             content_type="multipart/form-data")
    assert r.status_code == 302
    nama = q(client, "SELECT foto FROM transactions")[0]["foto"]
    tersimpan = Path(client.app.config["UPLOAD_FOLDER"]) / nama
    with Image.open(tersimpan) as gambar:
        assert max(gambar.size) <= 1600
    assert tersimpan.stat().st_size < 700_000
    assert tersimpan.stat().st_size < ukuran_asli * 0.25


def test_waktu_watermark_diambil_dari_metadata_foto_dan_fallback_browser():
    from app.photos import choose_photo_time

    foto, _ = jpeg_dengan_exif()
    with Image.open(foto) as gambar:
        waktu, dari_metadata = choose_photo_time(gambar, "2026-10-16 10:11:12", dt.datetime(2026, 10, 19, 9))
    assert dari_metadata is True
    assert waktu == dt.datetime(2026, 10, 15, 7, 8, 9)  # EXIF asli selalu diprioritaskan

    foto_tanpa_exif, _ = jpeg_file()
    with Image.open(foto_tanpa_exif) as gambar:
        waktu, dari_metadata = choose_photo_time(gambar, "2026-10-16 10:11:12", dt.datetime(2026, 10, 19, 9))
    assert dari_metadata is True
    assert waktu == dt.datetime(2026, 10, 16, 10, 11, 12)  # metadata yang dibaca sebelum kompresi browser


def test_pengaman_restart_dan_error(tmp_path, monkeypatch):
    """Server lama + kode baru: tampil peringatan restart; error tampil ramah & tercatat di data/error.log."""
    monkeypatch.setenv("STOK_DATA_DIR", str(tmp_path))
    app = create_app({"SECRET_KEY": "x"})  # tanpa TESTING agar errorhandler 500 aktif

    def boom():
        raise RuntimeError("uji error")
    app.add_url_rule("/_boom", "boom", boom)
    c = app.test_client()
    c.post("/login", data={"username": "admin", "password": "admin123"})
    assert "Aplikasi sudah diperbarui" not in c.get("/stok").data.decode()
    app.config["SIDIK_KODE"] = 0  # seolah file kode berubah setelah server start
    assert "Aplikasi sudah diperbarui" in c.get("/stok").data.decode()
    r = c.get("/_boom")
    assert r.status_code == 500 and "Data Anda tetap aman" in r.data.decode() and "start.bat" in r.data.decode()
    assert "uji error" in (tmp_path / "error.log").read_text(encoding="utf-8")

def test_pengembalian_barang_lama(client):
    """Tukar barang: barang lama dicatat & dilampirkan ke bon KELUAR, tidak mengubah stok."""
    upload_master(client, master_rows())
    ids = {r["material"] + r["special_stock"]: r["id"] for r in q(client, "SELECT id, material, special_stock FROM materials")}
    racun, egrek = ids["904.001.002"], ids["906.002.003"]
    foto = lambda: jpeg_file("x.jpg")  # noqa: E731
    r = post(client, "/master/atur", {"ids": [egrek], "wajib_kembali": "1"}, headers={"X-Requested-With": "fetch"})
    assert r.json["ok"]
    html = client.get("/master").data.decode()
    assert 'data-inline="wajib_kembali"' in html
    assert client.get("/api/material?q=egrek").json[0]["wajib_kembali"] == 1

    def keluar(extra, mats, qtys, kq, kk):
        return post(client, "/input", {"tombol": "keluar", "material_id": mats, "qty": qtys, "kembali_qty": kq,
                                       "kembali_kondisi": kk, "penerima": "Pemanen Andi", "foto": foto(), **extra},
                    content_type="multipart/form-data", follow_redirects=True).data.decode()

    # wajib tukar tapi barang lama tidak diisi -> ditolak
    assert "wajib tukar" in keluar({}, [egrek], ["2"], [""], [""])
    # barang lama melebihi qty keluar -> ditolak
    assert "melebihi qty keluar" in keluar({}, [egrek], ["2"], ["3"], ["Rusak"])
    # ada barang lama diserahkan tapi tanpa foto barang lama -> ditolak
    assert "Foto barang lama wajib" in keluar({}, [egrek], ["2"], ["2"], ["Rusak"])
    # egrek (wajib) + racun (tidak wajib) dalam satu bon
    html = keluar({"foto_kembali": jpeg_file("lama.jpg", (220, 230, 210))},
                  [egrek, racun], ["2", "1"], ["2", ""], ["Rusak", ""])
    assert "tersimpan, 2 barang" in html
    kb = q(client, "SELECT * FROM pengembalian")
    assert len(kb) == 1 and kb[0]["material_id"] == egrek and kb[0]["qty"] == 2 and kb[0]["kondisi"] == "Rusak"
    assert kb[0]["foto"] and kb[0]["dikembalikan_oleh"] == "Pemanen Andi"
    trx_egrek = q(client, "SELECT * FROM transactions WHERE id = ?", (kb[0]["transaksi_id"],))[0]
    assert trx_egrek["material_id"] == egrek and trx_egrek["bon"]
    assert stok(client, "906.002.003") == 8 - 2  # barang lama TIDAK menambah stok
    # hilang: boleh qty 0
    assert "tersimpan" in keluar({}, [egrek], ["1"], [""], ["Hilang / tidak dikembalikan"])
    assert q(client, "SELECT qty FROM pengembalian ORDER BY id DESC LIMIT 1")[0]["qty"] == 0
    # pekerja baru: tidak ada barang lama
    assert "tersimpan" in keluar({}, [egrek], ["1"], [""], ["Tidak ada — pekerja baru / tambahan"])
    assert q(client, "SELECT kondisi FROM pengembalian ORDER BY id DESC LIMIT 1")[0]["kondisi"].startswith("Tidak ada")

    # tampil di daftar transaksi
    halaman_transaksi = client.get(f"/transaksi?bon={trx_egrek['bon']}").data.decode()
    assert "2 Rusak" in halaman_transaksi and 'class="ikon"' in halaman_transaksi
    # racun dijadikan wajib belakangan -> keluar racun tadi masuk daftar belum kembali
    post(client, "/master/atur", {"ids": [racun], "wajib_kembali": "1"})
    halaman = client.get("/pengembalian").data.decode()
    assert "Belum mengembalikan barang lama (1)" in halaman and "RACUN ROUNDUP" in halaman
    assert "barang lama belum kembali" in client.get("/transaksi?status=belum_kembali").data.decode()
    assert "Barang lama belum kembali" in client.get("/").data.decode()
    trx_racun = q(client, "SELECT id FROM transactions WHERE material_id = ? AND jenis = 'keluar'", (racun,))[0]["id"]
    # catat susulan: tanpa foto ditolak, dengan foto tersimpan
    r = post(client, f"/transaksi/{trx_racun}/kembali", {"qty": "1", "kondisi": "Tabung / wadah kosong"},
             content_type="multipart/form-data", follow_redirects=True)
    assert "Foto barang lama wajib" in r.data.decode()
    r = post(client, f"/transaksi/{trx_racun}/kembali", {"qty": "1", "kondisi": "Tabung / wadah kosong",
                                                       "oleh": "Andi", "keterangan": "botol kosong",
                                                       "foto": foto()},
             content_type="multipart/form-data", follow_redirects=True)
    assert "tercatat" in r.data.decode()
    assert "Belum mengembalikan barang lama (0)" in client.get("/pengembalian").data.decode()
    # export & batal
    ws = openpyxl.load_workbook(io.BytesIO(client.get("/pengembalian/export").data)).active
    assert ws["I1"].value == "Kondisi" and ws.max_row == 1 + 4
    pid = q(client, "SELECT id FROM pengembalian WHERE material_id = ?", (racun,))[0]["id"]
    post(client, f"/pengembalian/{pid}/batal")
    assert "Belum mengembalikan barang lama (1)" in client.get("/pengembalian").data.decode()
    # kartu stok menampilkan info & tombol catat
    kartu = client.get(f"/stok/{egrek}").data.decode()
    assert "2 Rusak" in kartu and 'data-trx-dialog="dlgKembali"' in kartu and 'class="ikon"' in kartu


def test_upload_mulai_baru(client):
    """Mulai baru: file SAP jadi stok awal; transaksi belum diinput SAP tetap dihitung, yang sudah ✓ tidak dobel."""
    upload_master(client, master_rows())
    ids = {r["material"] + r["special_stock"]: r["id"] for r in q(client, "SELECT id, material, special_stock FROM materials")}
    racun, beras = ids["904.001.002"], ids["905.001.001"]
    foto = lambda: jpeg_file("x.jpg")  # noqa: E731
    for mid in (racun, beras):
        post(client, "/input", {"tombol": "keluar", "material_id": [mid], "qty": ["5"], "penerima": "A",
                                "foto": foto()}, content_type="multipart/form-data")
    # keluar beras sudah diinput ke SAP sebelum tarikan baru
    tb = q(client, "SELECT id FROM transactions WHERE material_id = ?", (beras,))[0]["id"]
    post(client, "/transaksi/posting", {"ids": [tb], "aksi": "posting"})
    post(client, "/master/atur", {"ids": [racun], "kategori_id": "1", "stok_minimum": "7"})
    # tarikan SAP baru: racun 60 (belum termasuk keluar 5), beras 35 (sudah termasuk keluar 5); egrek tidak ada
    rows = [r[:] for r in master_rows()]
    rows[1][7] = 60
    rows[2][7] = 35
    rows = [r for r in rows if r[2] != "906.002.003"]
    nanti = (dt.datetime.now() + dt.timedelta(minutes=5)).strftime("%Y-%m-%dT%H:%M")
    r = upload_master(client, rows, reset="1", sap_per=nanti)
    assert "Mulai baru" in r.data.decode()
    v = {x["material"] + x["special_stock"]: x for x in q(client, "SELECT * FROM v_stok")}
    assert v["904.001.002"]["stok"] == 55 and v["904.001.002"]["selisih_nyata"] == 0  # 60 - keluar belum di SAP
    assert v["904.001.002"]["kategori_id"] == 1 and v["904.001.002"]["stok_minimum"] == 7  # pengaturan tetap
    assert v["905.001.001"]["stok"] == 35 and v["905.001.001"]["selisih_nyata"] == 0  # tidak dikurangi dua kali
    assert v["906.002.003"]["stok"] == 0 and v["906.002.003"]["ada_di_sap"] == 0
    assert v["904.001.002"]["tgl_stok_awal"] == nanti[:10]


def test_minimum_cepat_dan_pr(client):
    upload_master(client, master_rows())
    v = lambda kode: q(client, "SELECT * FROM v_stok WHERE material = ? AND special_stock = ''", (kode,))[0]  # noqa: E731
    # isi minimum hanya untuk awalan 904
    r = post(client, "/master/minimum", {"aksi": "isi", "awalan": "904", "nilai": "150"}, follow_redirects=True)
    assert "untuk 2 material" in r.data.decode()  # racun biasa + racun Q
    assert v("904.001.002")["stok_minimum"] == 150 and v("905.001.001")["stok_minimum"] == 0
    # tanpa "timpa": yang sudah diisi tidak berubah
    post(client, "/master/minimum", {"aksi": "isi", "awalan": "", "nilai": "45"})
    assert v("904.001.002")["stok_minimum"] == 150 and v("905.001.001")["stok_minimum"] == 45
    post(client, "/master/minimum", {"aksi": "isi", "awalan": "906", "nilai": "3", "timpa": "1"})
    assert v("906.002.003")["stok_minimum"] == 3
    # reset per awalan, lalu isi lagi
    post(client, "/master/minimum", {"aksi": "reset", "awalan": "906"})
    assert v("906.002.003")["stok_minimum"] == 0 and v("905.001.001")["stok_minimum"] == 45

    # --- PR dari stok menipis: racun 100 <= 150, beras 40 <= 45
    r = post(client, "/pr/baru", {"judul": "PR UMUM Okt", "plant": "CM3A", "bulan_default": "3", "isi": "menipis"},
             follow_redirects=True)
    html = r.data.decode()
    assert "dengan 3 barang stok menipis" in html  # racun, racun Q, beras
    pr = q(client, "SELECT * FROM pr")[0]
    items = {i["material_id"]: i for i in q(client, "SELECT * FROM pr_item")}
    ids = {x["material"] + x["special_stock"]: x["id"] for x in q(client, "SELECT id, material, special_stock FROM materials")}
    racun = items[ids["904.001.002"]]
    assert racun["qty"] == 200 and racun["stok"] == 100 and racun["harga"] == 50000  # 2x150-100; 5jt/100
    # barang manual (kode belum ada di stok)
    r = post(client, f"/pr/{pr['id']}", {"aksi": "manual", "m_kode": "974.006.503", "m_nama": "Pen K-1 Kenko",
                                       "m_uom": "pac", "m_qty": "240", "m_harga": "4000"}, follow_redirects=True)
    assert "(manual) ditambahkan" in r.data.decode()
    man = q(client, "SELECT * FROM pr_item WHERE material_id IS NULL")[0]
    assert q(client, "SELECT judul FROM pr")[0]["judul"] == "PR UMUM Okt"  # judul tidak hilang saat tambah barang
    assert man["deskripsi"] == "PEN K-1 KENKO" and man["satuan"] == "PAC" and man["stok"] == 0
    # export sebelum ada nomor PR SAP (untuk attachment saat membuat PR di SAP): kolom PR kosong
    ws0 = openpyxl.load_workbook(io.BytesIO(client.get(f"/pr/{pr['id']}/export").data)).active
    assert ws0["D5"].value is None and ws0["E5"].value and ws0["T5"].value.startswith("=(G5+J5)")
    assert "Alur PR ke SAP" in client.get(f"/pr/{pr['id']}").data.decode()
    halaman = client.get(f"/pr/{pr['id']}").data.decode()
    assert "PR SAP per kelompok (3)" in halaman and all(g in halaman for g in ("Group 904", "Group 905", "Group 974"))
    # ubah qty racun + simpan PR (nomor SAP)
    data = {"aksi": "simpan", "tanggal": pr["tanggal"], "judul": "PR UMUM Okt", "nomor_sap": "1300008899", "status": "diajukan", "bulan_default": "3"}
    for it in q(client, "SELECT * FROM pr_item"):
        data.update({f"qty_{it['id']}": str(it["qty"]), f"harga_{it['id']}": str(it["harga"] or ""),
                     f"n_{it['id']}": "3", f"auto_{it['id']}": "1", f"periode_{it['id']}": "3 Bulan",
                     f"tp_{it['id']}": it["terakhir_purchase"] or ""})
    data[f"qty_{racun['id']}"] = "220"
    # nomor PR SAP per kelompok (default kelompok = Material Group karena kategori belum diatur)
    data.update({"sapk_1": "Group 904", "sapn_1": "1300008899", "sapk_2": "Group 974", "sapn_2": "1300008900"})
    data[f"nama_{man['id']}"] = "PEN, K-1 KENKO"
    data[f"kode_{man['id']}"] = "974.006.503"
    data[f"uom_{man['id']}"] = "pac"
    post(client, f"/pr/{pr['id']}", data)
    assert q(client, "SELECT qty FROM pr_item WHERE id = ?", (racun["id"],))[0]["qty"] == 220
    assert client.get("/pr/").status_code == 200 and client.get(f"/pr/{pr['id']}").status_code == 200

    # --- export Excel format Pendukung PR
    r = client.get(f"/pr/{pr['id']}/export")
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    assert ws["B2"].value.startswith("Pendukung PR") and ws["J3"].value == "Stock on Hand" and ws["M4"].value == "Bulan ke-1"
    baris = {ws[f"E{r}"].value: r for r in range(5, 5 + 4)}
    rr = next(r for r in range(5, 9) if ws[f"E{r}"].value == "904.001.002" and ws[f"J{r}"].value == 100)
    assert ws[f"C{rr}"].value == "CM3A" and ws[f"D{rr}"].value == 1300008899 and ws[f"G{rr}"].value == 220
    assert ws[f"H{rr}"].value == "L" and ws[f"I{rr}"].value == f"=G{rr}*50000" and ws[f"J{rr}"].value == 100
    assert [ws[f"{c}{rr}"].value for c in "MNO"] == [107, 107, 106] and ws[f"P{rr}"].value is None  # (220+100)/3
    assert ws[f"T{rr}"].value == f"=(G{rr}+J{rr})-SUM(M{rr}:R{rr})"
    rm = baris["974.006.503"]
    assert ws[f"D{rm}"].value == 1300008900  # nomor PR SAP kelompok lain
    assert ws[f"F{rm}"].value == "PEN, K-1 KENKO" and ws[f"J{rm}"].value == 0 and ws[f"I{rm}"].value == f"=G{rm}*4000"
    assert ws["F9"].value == "SUBTOTAL " and ws["I9"].value == "=SUM(I5:I8)"
    # export per kelompok (attachment satu PR SAP)
    ws2 = openpyxl.load_workbook(io.BytesIO(client.get(f"/pr/{pr['id']}/export?kelompok=Group 974").data)).active
    assert ws2["E5"].value == "974.006.503" and ws2["D5"].value == 1300008900 and ws2["F6"].value == "SUBTOTAL "
    ws3 = openpyxl.load_workbook(io.BytesIO(client.get(f"/pr/{pr['id']}/export?kelompok=Group 904").data)).active
    assert ws3["E5"].value == ws3["E6"].value == "904.001.002" and ws3["F7"].value == "SUBTOTAL "
    assert "1300008899, 1300008900" in client.get("/pr/").data.decode() or "1300008900, 1300008899" in client.get("/pr/").data.decode()
    # pindah kelompok satu barang
    data2 = dict(data, **{f"kel_{man['id']}": "ATK"})
    post(client, f"/pr/{pr['id']}", data2)
    assert q(client, "SELECT kelompok FROM pr_item WHERE id = ?", (man["id"],))[0]["kelompok"] == "ATK"
    # hapus PR
    post(client, f"/pr/{pr['id']}/hapus")
    assert not q(client, "SELECT * FROM pr") and not q(client, "SELECT * FROM pr_item")


def test_upload_hasil_sap_otomatis_memperbarui_progress_pr(client):
    upload_master(client, master_rows())
    post(client, "/pr/baru", {"judul": "Kebutuhan panen", "plant": "CM3A", "bulan_default": "3", "isi": ""})
    p = q(client, "SELECT * FROM pr")[0]
    material = {r["material"]: r["id"] for r in q(client, "SELECT id, material FROM materials WHERE special_stock = ''")}
    post(client, f"/pr/{p['id']}", {"aksi": "tambah", "material_id": material["904.001.002"]})
    post(client, f"/pr/{p['id']}", {"aksi": "tambah", "material_id": material["905.001.001"]})
    items = q(client, """SELECT i.qty, m.material, m.satuan FROM pr_item i
                           JOIN materials m ON m.id=i.material_id WHERE i.pr_id=? ORDER BY i.id""", (p["id"],))
    head = ["Date", "Plant", "PR", "Kode Material", "Material Description", "Qty PR", "UoM"]

    salah = [head,
             ["18.10.2026", "CM3A", 1300009101, items[0]["material"], "Racun", items[0]["qty"] + 1, items[0]["satuan"]],
             ["18.10.2026", "CM3A", 1300009102, items[1]["material"], "Beras", items[1]["qty"], items[1]["satuan"]]]
    r = post(client, "/pr/upload-sap", {"file": (xlsx(salah), f"Pendukung_PR_{p['nomor']}_uji.xlsx")},
             content_type="multipart/form-data", follow_redirects=True)
    assert "Upload PR SAP ditolak" in r.data.decode()
    assert q(client, "SELECT status FROM pr WHERE id=?", (p["id"],))[0]["status"] == "draft"
    assert not q(client, "SELECT * FROM pr_sap WHERE pr_id=?", (p["id"],))

    benar = [head,
             ["18.10.2026", "CM3A", 1300009101, items[0]["material"], "Racun", items[0]["qty"], items[0]["satuan"]],
             ["18.10.2026", "CM3A", 1300009102, items[1]["material"], "Beras", items[1]["qty"], items[1]["satuan"]]]
    r = post(client, "/pr/upload-sap", {"file": (xlsx(benar), f"Pendukung_PR_{p['nomor']}_uji.xlsx")},
             content_type="multipart/form-data", follow_redirects=True)
    html = r.data.decode()
    assert "diperbarui otomatis: 2 barang, 2 kelompok, 2 nomor PR SAP" in html
    assert "1300009101" in html and "1300009102" in html
    assert q(client, "SELECT status FROM pr WHERE id=?", (p["id"],))[0]["status"] == "diajukan"
    nomor = {r["kelompok"]: r["nomor_sap"] for r in q(client, "SELECT * FROM pr_sap WHERE pr_id=?", (p["id"],))}
    assert nomor == {"Group 904": "1300009101", "Group 905": "1300009102"}


def test_monitoring_penerimaan_pr_parsial_grn_dan_sisa(client):
    upload_master(client, master_rows())
    post(client, "/master/minimum", {"aksi": "isi", "awalan": "904", "nilai": "150", "timpa": "1"})
    post(client, "/master/minimum", {"aksi": "isi", "awalan": "905", "nilai": "45", "timpa": "1"})
    post(client, "/pr/baru", {"judul": "PR monitor GRN", "plant": "CM3A", "bulan_default": "3", "isi": ""})
    p = q(client, "SELECT * FROM pr")[0]
    material = {row["material"]: row["id"] for row in q(
        client, "SELECT id, material FROM materials WHERE special_stock = ''")}
    post(client, f"/pr/{p['id']}", {"aksi": "tambah", "material_id": material["904.001.002"]})
    post(client, f"/pr/{p['id']}", {"aksi": "tambah", "material_id": material["905.001.001"]})
    items = q(client, "SELECT * FROM pr_item WHERE pr_id = ? ORDER BY id", (p["id"],))
    racun, beras = items
    assert racun["qty"] == 200 and beras["qty"] == 50

    # GRN pertama: racun parsial, beras langsung lengkap.
    r = post(client, f"/pr/{p['id']}/penerimaan", {
        "grn": "grn-5001", "tanggal_terima": HARI_INI.isoformat(), "keterangan_terima": "Surat jalan 01",
        f"terima_{racun['id']}": "50", f"terima_{beras['id']}": "50",
    }, follow_redirects=True)
    html = r.data.decode()
    assert "GRN GRN-5001 tersimpan: 2 barang diterima" in html
    assert "datang sebagian" in html and "belum datang" in html
    assert len(q(client, "SELECT * FROM pr_penerimaan")) == 1
    assert [row["qty"] for row in q(client, "SELECT * FROM pr_penerimaan_item ORDER BY id")] == [50, 50]
    assert not q(client, "SELECT * FROM transactions")  # monitoring PR tidak menggandakan stok MASUK

    # Nomor GRN yang sama dan qty melebihi sisa ditolak.
    r = post(client, f"/pr/{p['id']}/penerimaan", {
        "grn": "GRN-5001", "tanggal_terima": HARI_INI.isoformat(), f"terima_{racun['id']}": "1",
    }, follow_redirects=True)
    assert "sudah tercatat pada PR ini" in r.data.decode()
    r = post(client, f"/pr/{p['id']}/penerimaan", {
        "grn": "GRN-5002", "tanggal_terima": HARI_INI.isoformat(), f"terima_{racun['id']}": "151",
    }, follow_redirects=True)
    assert "melebihi sisa PR 150" in r.data.decode()
    assert len(q(client, "SELECT * FROM pr_penerimaan")) == 1

    # GRN kedua melengkapi sisa; daftar PR ikut menunjukkan lengkap.
    r = post(client, f"/pr/{p['id']}/penerimaan", {
        "grn": "GRN-5002", "tanggal_terima": HARI_INI.isoformat(), f"terima_{racun['id']}": "150",
    }, follow_redirects=True)
    detail = r.data.decode()
    assert 'aria-valuenow="100"' in detail and ">2</b><small>barang lengkap" in detail
    daftar = client.get("/pr/").data.decode()
    assert "Lengkap 2/2" in daftar and "GRN terakhir" in daftar and "GRN-5002" in daftar

    # Item yang sudah diterima tidak boleh terhapus; GRN yang salah dapat dibatalkan.
    r = post(client, f"/pr/{p['id']}", {"aksi": f"hapus_{racun['id']}"}, follow_redirects=True)
    assert "sudah memiliki penerimaan GRN" in r.data.decode()
    rid2 = q(client, "SELECT id FROM pr_penerimaan WHERE grn = 'GRN-5002'")[0]["id"]
    r = post(client, f"/pr/{p['id']}/penerimaan/{rid2}/hapus", follow_redirects=True)
    assert "Qty kembali menjadi belum datang" in r.data.decode()
    assert "50" in r.data.decode() and len(q(client, "SELECT * FROM pr_penerimaan")) == 1

    # Hapus PR membersihkan riwayat GRN terkait tanpa meninggalkan data yatim.
    post(client, f"/pr/{p['id']}/hapus")
    assert not q(client, "SELECT * FROM pr_penerimaan")
    assert not q(client, "SELECT * FROM pr_penerimaan_item")


def test_opname_banyak_barang_dan_batal_sesi(client):
    """Semua barang yang diisi tersimpan dalam satu sesi; sesi yang salah bisa dibatalkan & stok kembali."""
    upload_master(client, master_rows())
    ids = {r["material"] + r["special_stock"]: r["id"] for r in q(client, "SELECT id, material, special_stock FROM materials")}
    racun, beras, egrek = ids["904.001.002"], ids["905.001.001"], ids["906.002.003"]
    # barang dari "pencarian berbeda" tetap terkirim bersama (draft di browser menambah input ke form)
    r = post(client, "/opname", {f"fisik_{racun}": "98", f"fisik_{beras}": "40", f"fisik_{egrek}": "0",
                                 "jenis_barang": "LAINNYA", "alasan": "Stock opname (hitung fisik)"}, follow_redirects=True)
    assert "3 material dihitung, 2 disesuaikan" in r.data.decode() and "opnameTersimpan()" in r.data.decode()
    assert stok(client, "906.002.003") == 0 and stok(client, "904.001.002") == 98
    sid = q(client, "SELECT id FROM opname_sesi")[0]["id"]
    # halaman opname memuat data tiap baris untuk draft
    assert 'data-opname-row data-mid=' in client.get("/opname").data.decode()
    # batalkan sesi -> penyesuaian dibatalkan, stok kembali, sesi hilang dari riwayat
    r = post(client, f"/opname/{sid}/batal", follow_redirects=True)
    assert "2 penyesuaian stok dikembalikan" in r.data.decode()
    assert stok(client, "906.002.003") == 8 and stok(client, "904.001.002") == 100
    assert not q(client, "SELECT * FROM opname_sesi") and not q(client, "SELECT * FROM opname_item")
    assert len(q(client, "SELECT * FROM transactions WHERE jenis = 'penyesuaian' AND batal = 1")) == 2


def test_opname_bersama_beberapa_petugas_satu_gudang(client):
    """Admin membuka sesi; petugas menghitung paralel tanpa saling menimpa; stok diposting saat ditutup."""
    upload_master(client, master_rows())
    ids = {r["material"] + r["special_stock"]: r["id"]
           for r in q(client, "SELECT id, material, special_stock FROM materials")}
    racun, beras = ids["904.001.002"], ids["905.001.001"]
    for no in (1, 2):
        post(client, "/pengguna", {"username": f"hitung{no}", "nama": f"Penghitung {no}",
                                    "password": "rahasia1", "role": "petugas", "kebun_kode": "CM3A"})
    r = post(client, "/opname/mulai", {"plant": "CM3A", "jenis_barang": "SPAREPART",
                                        "tanggal": HARI_INI.isoformat(), "keterangan": "hitung bersama"})
    assert r.status_code == 302 and "/hitung" in r.headers["Location"]
    sid = q(client, "SELECT id FROM opname_sesi WHERE status = 'terbuka'")[0]["id"]

    def petugas(username):
        c = client.app.test_client()
        c.post("/login", data={"username": username, "password": "rahasia1"})
        c.get("/opname")
        with c.session_transaction() as s:
            c.csrf = s["csrf"]
        return c

    p1, p2 = petugas("hitung1"), petugas("hitung2")
    assert b"Sesi yang sedang berjalan" in p1.get("/opname").data
    r = post(p1, f"/opname/{sid}/hitung", {f"fisik_{racun}": "98", f"ket_{racun}": "rak A"},
             follow_redirects=True)
    assert b"hitungan Anda tersimpan" in r.data
    assert b'id="opnameCariCepat"' in r.data and b"kartu-opname" in r.data and b"data-search=" in r.data
    assert b'id="opnamePrioritas"' in r.data and b'id="opnameBaris"' in r.data
    r = post(p2, f"/opname/{sid}/hitung", {f"fisik_{beras}": "39", f"fisik_{racun}": "97"},
             follow_redirects=True)
    assert b"sudah dihitung oleh Penghitung 1" in r.data
    item = q(client, "SELECT fisik FROM opname_item WHERE sesi_id = ? AND material_id = ?", (sid, racun))[0]
    assert item["fisik"] == 98  # hitungan petugas pertama tidak tertimpa
    assert len(q(client, "SELECT DISTINCT counted_by FROM opname_item WHERE sesi_id = ?", (sid,))) == 2

    # Selama sesi terbuka transaksi normal gudang yang sama ditahan dan stok belum disesuaikan.
    r = post(p1, "/input", {"tombol": "masuk", "material_id": racun, "qty": "1"},
             content_type="multipart/form-data", follow_redirects=True)
    assert b"sedang Stock Opname bersama" in r.data
    assert stok(client, "904.001.002") == 100 and stok(client, "905.001.001") == 40

    r = post(client, f"/opname/{sid}/selesai", {"konfirmasi_semua": "1"}, follow_redirects=True)
    assert b"2 material dihitung, 2 disesuaikan" in r.data
    assert stok(client, "904.001.002") == 98 and stok(client, "905.001.001") == 39
    sesi = q(client, "SELECT * FROM opname_sesi WHERE id = ?", (sid,))[0]
    assert sesi["status"] == "selesai" and sesi["closed_by"] is not None and sesi["closed_at"]


def test_admin_dapat_mendelegasikan_pembukaan_sesi_opname_sementara(client):
    """Delegasi membuka/menutup sesi sendiri dengan konfirmasi semua anggota sudah menyimpan."""
    upload_master(client, master_rows())
    post(client, "/pengguna", {"username": "wakilopname", "nama": "Wakil Opname",
                                "password": "rahasia1", "role": "petugas", "kebun_kode": "CM3A"})
    petugas_row = q(client, "SELECT * FROM users WHERE username = 'wakilopname'")[0]

    petugas = client.app.test_client()
    petugas.post("/login", data={"username": "wakilopname", "password": "rahasia1"})
    petugas.get("/opname")
    with petugas.session_transaction() as sesi_login:
        petugas.csrf = sesi_login["csrf"]

    # Tanpa delegasi, petugas tetap tidak dapat membuka sesi melalui request langsung.
    r = post(petugas, "/opname/mulai", {"plant": "CM3A", "jenis_barang": "RACUN"},
             follow_redirects=True)
    assert b"tidak tersedia atau sudah kedaluwarsa" in r.data
    assert not q(client, "SELECT * FROM opname_sesi")

    berlaku = (dt.datetime.now() + dt.timedelta(hours=2)).strftime("%Y-%m-%dT%H:%M")
    r = post(client, "/opname/delegasi", {
        "aksi": "beri", "user_id": petugas_row["id"], "valid_until": berlaku,
    }, follow_redirects=True)
    assert b"dapat membuka sesi opname sampai" in r.data
    delegasi = q(client, "SELECT * FROM opname_delegasi WHERE user_id = ?", (petugas_row["id"],))[0]
    assert delegasi["plant"] == "CM3A" and delegasi["granted_by"] is not None
    assert b"Delegasi aktif" in petugas.get("/opname").data

    # Plant dari formulir tidak dipercaya; sesi selalu mengikuti gudang akun petugas.
    r = post(petugas, "/opname/mulai", {
        "plant": "CM9A", "jenis_barang": "RACUN", "tanggal": HARI_INI.isoformat(),
        "keterangan": "admin sedang di lapangan",
    })
    assert r.status_code == 302 and "/hitung" in r.headers["Location"]
    sesi = q(client, "SELECT * FROM opname_sesi WHERE status = 'terbuka'")[0]
    assert sesi["plant"] == "CM3A" and sesi["created_by"] == petugas_row["id"]
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    post(petugas, f"/opname/{sesi['id']}/hitung", {f"fisik_{racun}": "98"})
    halaman = petugas.get(f"/opname/{sesi['id']}/hitung").data
    assert b"Menutup sesi akan menyelesaikan Stock Opname" in halaman
    assert b"Wakil Opname" in halaman and b"1 barang" in halaman

    # Tanpa konfirmasi eksplisit sesi tidak ditutup, sehingga draft anggota lain tidak terlewat.
    r = post(petugas, f"/opname/{sesi['id']}/selesai", follow_redirects=True)
    assert b"Pastikan semua anggota sudah menyimpan hitungannya" in r.data
    assert q(client, "SELECT status FROM opname_sesi WHERE id = ?", (sesi["id"],))[0]["status"] == "terbuka"

    r = post(petugas, f"/opname/{sesi['id']}/selesai", {"konfirmasi_semua": "1"},
             follow_redirects=True)
    assert b"1 material dihitung, 1 disesuaikan" in r.data
    selesai = q(client, "SELECT * FROM opname_sesi WHERE id = ?", (sesi["id"],))[0]
    assert selesai["status"] == "selesai" and selesai["closed_by"] == petugas_row["id"]
    assert stok(client, "904.001.002") == 98

    # Admin tetap dapat mencabut delegasi kapan saja.
    r = post(client, "/opname/delegasi", {
        "aksi": "cabut", "user_id": petugas_row["id"],
    }, follow_redirects=True)
    assert b"Delegasi opname untuk Wakil Opname dicabut" in r.data
    assert not q(client, "SELECT * FROM opname_delegasi WHERE user_id = ?", (petugas_row["id"],))
    r = post(petugas, "/opname/mulai", {"plant": "CM3A", "jenis_barang": "BERAS"},
             follow_redirects=True)
    assert b"tidak tersedia atau sudah kedaluwarsa" in r.data
    aksi = [row["aksi"] for row in q(client, "SELECT aksi FROM audit_log")]
    assert "beri_delegasi_opname" in aksi and "cabut_delegasi_opname" in aksi


def test_api_mobile_sinkron_idempoten_duplikat_dan_perangkat_admin(client):
    """APK dapat sinkron, retry tidak menggandakan BON, dan admin dapat mencabut perangkat."""
    preflight = client.options("/api/mobile/v1/bootstrap", headers={
        "Origin": "http://localhost",
        "Access-Control-Request-Method": "GET",
        "Access-Control-Request-Headers": "authorization,x-app-version",
    })
    assert preflight.status_code in {200, 204}
    assert preflight.headers["Access-Control-Allow-Origin"] == "http://localhost"
    allowed_headers = preflight.headers["Access-Control-Allow-Headers"].lower()
    assert "authorization" in allowed_headers and "x-app-version" in allowed_headers

    upload_master(client, master_rows(qty_racun=20))
    post(client, "/aplikasi-mobile", {
        "action": "save_settings",
        "github_repo": "https://github.com/contoh/monitoring-stok.git",
        "minimum_version": "0.1.0",
    })
    device_uuid = str(uuid.uuid4())
    headers = mobile_login(client, device_uuid)

    bootstrap = client.get("/api/mobile/v1/bootstrap", headers=headers)
    assert bootstrap.status_code == 200
    data = bootstrap.get_json()
    assert data["update"]["github_repo"] == "contoh/monitoring-stok"
    assert data["valid_for_date"] == HARI_INI.isoformat()
    racun = next(m for m in data["materials"] if m["material"] == "904.001.002" and m["stok"] == 20)

    client_uuid = str(uuid.uuid4())
    payload = {
        "client_uuid": client_uuid,
        "type": "keluar",
        "date": HARI_INI.isoformat(),
        "recipient": "Bobi",
        "note": "Afd 3",
        "target": {"type": "blok", "code": "D24", "activity": ""},
        "items": [{"material_id": racun["id"], "qty": 2}],
    }
    first = client.post("/api/mobile/v1/transactions", headers=headers, data={
        "payload": json.dumps(payload), "photo": jpeg_file("apk-1.jpg"),
    }, content_type="multipart/form-data")
    assert first.status_code == 200
    bon = first.get_json()["bon"]
    assert bon.startswith("BON-") and first.get_json()["already_received"] is False
    assert stok(client, "904.001.002") == 18

    retry = client.post("/api/mobile/v1/transactions", headers=headers, data={
        "payload": json.dumps(payload), "photo": jpeg_file("apk-retry.jpg"),
    }, content_type="multipart/form-data")
    assert retry.status_code == 200
    assert retry.get_json()["already_received"] is True and retry.get_json()["bon"] == bon
    assert stok(client, "904.001.002") == 18
    assert len(q(client, "SELECT * FROM mobile_submissions")) == 1

    duplicate_payload = {**payload, "client_uuid": str(uuid.uuid4())}
    duplicate = client.post("/api/mobile/v1/transactions", headers=headers, data={
        "payload": json.dumps(duplicate_payload), "photo": jpeg_file("apk-duplikat.jpg"),
    }, content_type="multipart/form-data")
    assert duplicate.status_code == 409
    assert duplicate.get_json()["duplicate"] is True and duplicate.get_json()["matches"][0]["bon"] == bon

    duplicate_payload["confirm_duplicate"] = True
    confirmed = client.post("/api/mobile/v1/transactions", headers=headers, data={
        "payload": json.dumps(duplicate_payload), "photo": jpeg_file("apk-benar.jpg"),
    }, content_type="multipart/form-data")
    assert confirmed.status_code == 200 and confirmed.get_json()["bon"] != bon
    assert stok(client, "904.001.002") == 16
    assert len(q(client, "SELECT * FROM transactions WHERE sumber='apk' AND batal=0")) == 2

    device = q(client, "SELECT * FROM mobile_devices WHERE device_uuid=?", (device_uuid,))[0]
    page = client.get("/aplikasi-mobile").data.decode()
    assert "HP Uji Gudang" in page and "contoh/monitoring-stok" in page
    post(client, "/aplikasi-mobile", {"action": "disable_device", "device_id": device["id"]})
    denied = client.get("/api/mobile/v1/bootstrap", headers=headers)
    assert denied.status_code == 401
    relogin = client.post("/api/mobile/v1/login", json={
        "username": "admin", "password": "admin123", "device_uuid": device_uuid,
        "device_name": "HP Uji Gudang", "app_version": "0.1.0",
    })
    assert relogin.status_code == 403


def test_api_mobile_menolak_material_dikunci_dan_tanpa_foto(client):
    """Validasi kritis server tetap berlaku walaupun payload berasal dari APK."""
    upload_master(client, master_rows(qty_racun=10))
    headers = mobile_login(client)
    racun = q(client, "SELECT id FROM materials WHERE material='904.001.002' AND special_stock='' ")[0]["id"]
    payload = {
        "client_uuid": str(uuid.uuid4()), "type": "keluar", "date": HARI_INI.isoformat(),
        "recipient": "Bobi", "target": {"type": "blok", "code": "D24"},
        "items": [{"material_id": racun, "qty": 1}],
    }
    no_photo = client.post("/api/mobile/v1/transactions", headers=headers,
                           data={"payload": json.dumps(payload)}, content_type="multipart/form-data")
    assert no_photo.status_code == 400 and "Foto orang" in no_photo.get_json()["error"]

    import sqlite3
    db = sqlite3.connect(client.app.config["DATABASE"])
    db.execute("UPDATE materials SET input_locked=1, input_lock_reason='Gunakan kode pengganti' WHERE id=?", (racun,))
    db.commit()
    db.close()
    locked = client.post("/api/mobile/v1/transactions", headers=headers, data={
        "payload": json.dumps(payload), "photo": jpeg_file("terkunci.jpg"),
    }, content_type="multipart/form-data")
    assert locked.status_code == 409 and "Gunakan kode pengganti" in locked.get_json()["error"]
    assert not q(client, "SELECT * FROM mobile_submissions")
