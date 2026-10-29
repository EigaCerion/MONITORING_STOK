"""Baca file Excel: export master SAP dan template transaksi masuk/transfer."""
import datetime as dt
import re

import openpyxl

MASTER_COLS = {
    "name 1": "name1",
    "material group": "material_group",
    "material": "material",
    "material description": "deskripsi",
    "plant": "plant",
    "storage location": "storage_location",
    "batch": "batch",
    "unrestricted": "unrestricted",
    "base unit of measure": "satuan",
    "value unrestricted": "value_unrestricted",
    "special stock": "special_stock",
    "special stock number": "special_stock_no",
    "stock in transit": "stock_in_transit",
    "material type": "material_type",
}
MASTER_WAJIB = ["material", "deskripsi", "storage_location", "unrestricted", "satuan"]
MASTER_NUM = ("unrestricted", "value_unrestricted", "stock_in_transit")
GUDANG_SLOC = ("ST01", "STTF")

TRX_COLS = {
    "tanggal": "tanggal",
    "jenis transaksi": "jenis",
    "material": "material",
    "quantity": "qty",
    "base unit of measure": "satuan",
    "nomor dokumen": "nomor_dokumen",
    "nama penerima": "penerima",
    "kebun asal": "kebun_asal",
    "kebun tujuan": "kebun_tujuan",
    "keterangan": "keterangan",
}
TRX_WAJIB = ["tanggal", "jenis", "material", "qty"]
JENIS_MAP = {
    "MASUK": "masuk",
    "TRANSFER_MASUK": "transfer_masuk",
    "TRANSFER_KELUAR": "transfer_keluar",
    "KELUAR": "keluar",
}

MB51_COLS = {
    "material document": "sap_doc",
    "posting date": "tanggal",
    "wbs element": "special_stock_no",
    "material": "material",
    "plant": "plant",
    "storage location": "storage_location",
    "movement type": "movement_type",
    "quantity": "qty",
    "material description": "deskripsi",
    "special stock": "special_stock",
    "text": "teks_sap",
    "purchase order": "purchase_order",
}
MB51_WAJIB = ["sap_doc", "tanggal", "material", "plant", "movement_type", "qty"]

PR_PROGRESS_COLS = {
    "date": "tanggal",
    "plant": "plant",
    "pr": "sap_pr",
    "kode material": "material",
    "qty pr": "qty",
    "uom": "satuan",
    "material description": "deskripsi",
}
PR_PROGRESS_WAJIB = ["plant", "sap_pr", "material", "qty"]


class ExcelError(ValueError):
    pass


def teks(v):
    """Nilai sel -> string rapi (angka bulat tanpa '.0')."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def angka(v):
    """Nilai sel -> float. Menerima angka atau teks format 1.234,5 / 1234.5. None jika kosong/tidak valid."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    neg = s.endswith("-")  # format SAP: 10-
    s = s.rstrip("-")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        x = float(s)
    except ValueError:
        return None
    return -x if neg else x


def tanggal(v):
    """Nilai sel -> 'YYYY-MM-DD' atau None."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # serial Excel
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    s = str(v).strip()
    for f in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return dt.datetime.strptime(s, f).date().isoformat()
        except ValueError:
            pass
    return None


def _buka(path):
    """Buka workbook Excel dan ubah kegagalan parser menjadi pesan ramah pengguna."""
    try:
        return openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001 - file rusak / bukan xlsx
        raise ExcelError(f"File tidak bisa dibaca sebagai Excel (.xlsx): {e}") from e


def _cari_header(ws, kolom, wajib, prefix_ok=("tanggal",)):
    """Temukan baris header dan petakan nama kolom berdasarkan alias yang dikenal."""
    rows = list(ws.iter_rows(values_only=True))
    for i, row in enumerate(rows[:15]):
        peta = {}
        for j, v in enumerate(row):
            h = re.sub(r"\s+", " ", teks(v).lower())
            if h in kolom:
                peta.setdefault(kolom[h], j)
            else:
                for p in prefix_ok:
                    if h.startswith(p) and p in kolom:
                        peta.setdefault(kolom[p], j)
        if all(k in peta for k in wajib):
            return i, peta, rows
    nama = [k for k, v in kolom.items() if v in wajib]
    raise ExcelError("Header tidak ditemukan. Kolom wajib: " + ", ".join(n.title() for n in nama))


def baca_master_sap(path):
    """Baca saldo gudang ST01 + STTF; storage location lain tidak masuk stok aplikasi."""
    wb = _buka(path)
    ws = wb.worksheets[0]
    hi, peta, rows = _cari_header(ws, MASTER_COLS, MASTER_WAJIB)
    hasil, peringatan, lewati, lokasi_diabaikan = {}, [], 0, {}
    for n, row in enumerate(rows[hi + 1:], start=hi + 2):
        def ambil(k):
            """Ambil nilai kolom logis dari baris master saat ini."""
            j = peta.get(k)
            return row[j] if j is not None and j < len(row) else None

        material = teks(ambil("material"))
        if not material or material.lower().startswith("catatan"):
            lewati += 1
            continue
        d = {k: teks(ambil(k)) for k in MASTER_COLS.values() if k not in MASTER_NUM}
        sloc = d["storage_location"].upper()
        if sloc not in GUDANG_SLOC:
            lewati += 1
            label = sloc or "(kosong)"
            lokasi_diabaikan[label] = lokasi_diabaikan.get(label, 0) + 1
            continue
        d["storage_location"] = sloc
        for k in MASTER_NUM:
            x = angka(ambil(k))
            if x is None and k == "unrestricted" and teks(ambil(k)):
                peringatan.append(f"Baris {n}: Unrestricted '{ambil(k)}' bukan angka, dianggap 0")
            d[k] = x or 0.0
        d["special_stock"] = d["special_stock"].upper()
        kunci = (d["plant"], material, d["special_stock"], d["special_stock_no"])
        if kunci in hasil:
            lama = hasil[kunci]
            for k in MASTER_NUM:
                lama[k] += d[k]
            lama["_gabung"] += 1
            lama["_slocs"].add(sloc)
        else:
            d["_baris"] = n
            d["_gabung"] = 0
            d["_slocs"] = {sloc}
            hasil[kunci] = d
    wb.close()
    for d in hasil.values():
        lokasi = d.pop("_slocs")
        d["storage_location"] = "+".join(s for s in GUDANG_SLOC if s in lokasi)
    if lokasi_diabaikan:
        rincian = ", ".join(f"{nama}: {jumlah}" for nama, jumlah in sorted(lokasi_diabaikan.items()))
        peringatan.append(
            f"{sum(lokasi_diabaikan.values())} baris di luar ST01/STTF tidak dihitung sebagai stok gudang ({rincian})"
        )
    gabung = [d for d in hasil.values() if d["_gabung"]]
    if gabung:
        peringatan.append(
            f"{len(gabung)} material muncul di beberapa baris (beda storage location/batch) dan qty-nya dijumlahkan, "
            f"mis. {gabung[0]['material']}"
        )
    if not hasil:
        raise ExcelError("Tidak ada baris material yang terbaca. Pastikan data dimulai di bawah header.")
    return list(hasil.values()), peringatan


def baca_supplementary(path):
    """Baca master blok dan unit/VRA dari workbook Supplementary Plant Maintenance."""
    wb = _buka(path)

    def sheet_rows(sheet_name, required):
        """Temukan header bernama sama pada 10 baris pertama lalu kembalikan data terpetakan."""
        if sheet_name not in wb.sheetnames:
            return []
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        for header_index, row in enumerate(rows[:10]):
            mapping = {}
            for index, value in enumerate(row):
                header = re.sub(r"\s+", " ", teks(value).lower())
                if header:
                    # VRA ORDER memiliki dua kolom bernama "Description". Kolom
                    # pertama (setelah Equipment) berisi kode DT/EX/TR yang dicari
                    # operator; kolom kedua hanya deskripsi order SAP.
                    mapping.setdefault(header, index)
            if all(name in mapping for name in required):
                return [(number, values, mapping)
                        for number, values in enumerate(rows[header_index + 1:], start=header_index + 2)]
        raise ExcelError(f"Sheet {sheet_name}: header wajib tidak ditemukan ({', '.join(required)})")

    hasil, peringatan = {}, []
    try:
        block_rows = sheet_rows("BLOCK", ("estate", "division", "block", "block name", "phase"))
        if not block_rows:
            raise ExcelError("Sheet BLOCK tidak ditemukan pada file Supplementary")
        for number, row, mapping in block_rows:
            ambil = lambda name: row[mapping[name]] if mapping[name] < len(row) else None  # noqa: E731
            plant = teks(ambil("estate")).upper()
            code = teks(ambil("block name")).upper()
            sap_code = teks(ambil("block")).upper()
            if not plant or not code or not sap_code:
                continue
            key = (plant, "blok", code)
            hasil[key] = {
                "plant": plant, "target_type": "blok", "code": code, "sap_code": sap_code,
                "name": f"Blok {code}", "division": teks(ambil("division")),
                "phase": teks(ambil("phase")).upper(), "source": f"BLOCK baris {number}",
            }

        vra_rows = sheet_rows("VRA ORDER", ("order", "equipment", "description", "planning plant"))
        for number, row, mapping in vra_rows:
            ambil = lambda name: row[mapping[name]] if mapping[name] < len(row) else None  # noqa: E731
            plant = teks(ambil("planning plant")).upper()
            description = teks(ambil("description"))
            sap_code = teks(ambil("order")).upper()
            equipment = teks(ambil("equipment")).upper()
            code_match = re.match(r"\s*([A-Z]{1,5}\d{1,4})\s*(?:-|$)", description.upper())
            code = code_match.group(1) if code_match else equipment
            if not plant or not code or not sap_code:
                continue
            key = (plant, "unit", code)
            hasil[key] = {
                "plant": plant, "target_type": "unit", "code": code, "sap_code": sap_code,
                "name": description or equipment, "division": "", "phase": "",
                "source": f"VRA ORDER baris {number}",
            }
        if not vra_rows:
            peringatan.append("Sheet VRA ORDER tidak ada; hanya master blok yang diperbarui")
    finally:
        wb.close()
    if not hasil:
        raise ExcelError("Tidak ada blok atau unit yang terbaca dari file Supplementary")
    return list(hasil.values()), peringatan


def baca_transaksi(path):
    """Kembalikan list dict baris mentah dengan nomor baris Excel (_baris)."""
    wb = _buka(path)
    ws = wb.worksheets[0]
    hi, peta, rows = _cari_header(ws, TRX_COLS, TRX_WAJIB)
    hasil = []
    for n, row in enumerate(rows[hi + 1:], start=hi + 2):
        def ambil(k):
            """Ambil nilai kolom logis dari baris transaksi saat ini."""
            j = peta.get(k)
            return row[j] if j is not None and j < len(row) else None

        if all(v is None or teks(v) == "" for v in row):
            continue
        pertama = teks(row[0]) if row else ""
        if pertama.lower().startswith(("jenis:", "catatan")) and not teks(ambil("material")):
            continue
        jenis_raw = teks(ambil("jenis")).upper().replace(" ", "_")
        hasil.append({
            "_baris": n,
            "tanggal_raw": teks(ambil("tanggal")),
            "tanggal": tanggal(ambil("tanggal")),
            "jenis_raw": jenis_raw,
            "jenis": JENIS_MAP.get(jenis_raw),
            "material": teks(ambil("material")),
            "qty_raw": teks(ambil("qty")),
            "qty": angka(ambil("qty")),
            "satuan": teks(ambil("satuan")),
            "nomor_dokumen": teks(ambil("nomor_dokumen")),
            "penerima": teks(ambil("penerima")),
            "kebun_asal_raw": teks(ambil("kebun_asal")),
            "kebun_tujuan_raw": teks(ambil("kebun_tujuan")),
            "keterangan": teks(ambil("keterangan")),
        })
    wb.close()
    if not hasil:
        raise ExcelError("Tidak ada baris transaksi yang terbaca.")
    return hasil


def baca_mb51(path):
    """Baca baris material dari export SAP MB51; subtotal dan baris jasa otomatis dilewati."""
    wb = _buka(path)
    ws = wb.worksheets[0]
    hi, peta, rows = _cari_header(ws, MB51_COLS, MB51_WAJIB, prefix_ok=())
    hasil, dilewati = [], 0
    for n, row in enumerate(rows[hi + 1:], start=hi + 2):
        def ambil(k):
            """Ambil nilai kolom logis dari baris MB51 saat ini."""
            j = peta.get(k)
            return row[j] if j is not None and j < len(row) else None

        sap_doc = teks(ambil("sap_doc"))
        material = teks(ambil("material")).upper()
        plant = teks(ambil("plant")).upper()
        movement_type = teks(ambil("movement_type"))
        qty = angka(ambil("qty"))
        tgl = tanggal(ambil("tanggal"))
        # Export MB51 dapat memuat baris jasa (tanpa material) dan baris subtotal.
        if not all((sap_doc, material, plant, movement_type, tgl)) or qty is None or abs(qty) < 1e-9:
            dilewati += 1
            continue
        special_stock = teks(ambil("special_stock")).upper()
        # WBS Element pada MB51 juga terisi untuk pembebanan biaya transaksi
        # normal (mis. MvT 221). Ia baru menjadi identitas stok khusus bila
        # indikator Special Stock memang terisi. Tanpa aturan ini, transaksi
        # ST01 biasa gagal dipasangkan dengan material unrestricted aplikasi.
        special_stock_no = teks(ambil("special_stock_no")) if special_stock else ""
        hasil.append({
            "_baris": n,
            "sap_doc": sap_doc,
            "tanggal": tgl,
            "material": material,
            "plant": plant,
            "storage_location": teks(ambil("storage_location")).upper(),
            "movement_type": movement_type,
            "qty": round(qty, 3),
            "deskripsi": teks(ambil("deskripsi")),
            "special_stock": special_stock,
            "special_stock_no": special_stock_no,
            "teks_sap": teks(ambil("teks_sap")),
            "purchase_order": teks(ambil("purchase_order")),
        })
    wb.close()
    if not hasil:
        raise ExcelError("Tidak ada transaksi material yang terbaca dari MB51.")
    return hasil, dilewati


def baca_progress_pr(path):
    """Baca file Pendukung PR yang sudah diberi nomor PR oleh SAP."""
    wb = _buka(path)
    ws = wb.worksheets[0]
    hi, peta, rows = _cari_header(ws, PR_PROGRESS_COLS, PR_PROGRESS_WAJIB, prefix_ok=())
    hasil, error = [], []
    for n, row in enumerate(rows[hi + 1:], start=hi + 2):
        def ambil(k):
            """Ambil nilai kolom logis dari baris Progress PR saat ini."""
            j = peta.get(k)
            return row[j] if j is not None and j < len(row) else None

        material = teks(ambil("material")).upper()
        if not material:  # subtotal dan baris catatan
            continue
        sap_pr = teks(ambil("sap_pr"))
        qty = angka(ambil("qty"))
        plant = teks(ambil("plant")).upper()
        if not re.fullmatch(r"\d{8,12}", sap_pr):
            error.append(f"baris {n}: nomor PR SAP belum diisi atau tidak valid")
        if not plant:
            error.append(f"baris {n}: Plant kosong")
        if qty is None or qty < 0:
            error.append(f"baris {n}: Qty PR tidak valid")
        hasil.append({
            "_baris": n,
            "tanggal": tanggal(ambil("tanggal")),
            "plant": plant,
            "sap_pr": sap_pr,
            "material": material,
            "qty": qty,
            "satuan": teks(ambil("satuan")).upper(),
            "deskripsi": teks(ambil("deskripsi")),
        })
    wb.close()
    if not hasil:
        raise ExcelError("Tidak ada baris material yang terbaca dari file Pendukung PR.")
    if error:
        raise ExcelError("; ".join(error[:8]) + ("; dan lainnya" if len(error) > 8 else ""))
    return hasil
