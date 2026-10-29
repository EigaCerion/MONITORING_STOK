"""PR (Purchase Request) pemesanan barang ke pusat: saran qty, harga, pembagian periode, export Excel
"Pendukung PR (Stock, Terakhir GRN, Periode Pemakaian dan Estimasi GI)"."""
import datetime as dt
import io
import math
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .report import harga_satuan

STATUS = {"draft": "Draft", "diajukan": "Diajukan ke pusat", "selesai": "Selesai / barang datang", "batal": "Batal"}
TIPIS = Side(style="thin", color="000000")
KOTAK = Border(left=TIPIS, right=TIPIS, top=TIPIS, bottom=TIPIS)
TENGAH = Alignment(horizontal="center", vertical="center", wrap_text=True)
KEPALA = PatternFill("solid", fgColor="D9E1F2")


def harga_default(m):
    """Hitung harga satuan estimasi dari nilai dan qty unrestricted SAP."""
    h = harga_satuan(m)
    return round(h, 2) if h else None


def saran_qty(stok, minimum):
    """Pesan hingga 2x stok minimum (minimal sebesar stok minimum). 0 bila minimum belum diatur."""
    if not minimum or minimum <= 0:
        return 0.0
    return float(max(math.ceil(2 * minimum - max(stok, 0)), math.ceil(minimum)))


def terakhir_masuk(db, material_id):
    """Cari tanggal penerimaan terakhir material untuk referensi PR."""
    r = db.execute("""SELECT MAX(tanggal) FROM transactions WHERE material_id = ? AND batal = 0
                      AND jenis IN ('masuk', 'transfer_masuk')""", (material_id,)).fetchone()[0]
    return dt.date.fromisoformat(r).strftime("%d.%m.%Y") if r else "Belum ada data GRN"


def pakai_per_bulan(db, material_id, hari=90):
    """Hitung rata-rata pemakaian bulanan berdasarkan transaksi keluar terbaru."""
    awal = (dt.date.today() - dt.timedelta(days=hari)).isoformat()
    total = db.execute("""SELECT COALESCE(SUM(qty), 0) FROM transactions WHERE material_id = ? AND batal = 0
                          AND jenis IN ('keluar', 'transfer_keluar') AND tanggal >= ?""",
                       (material_id, awal)).fetchone()[0]
    return round(total / (hari / 30), 2)


def bagi_rata(total, n):
    """Bagi (Qty PR + Stock on Hand) ke n bulan (maks 6) supaya crosscek = 0. Sisa pembulatan ke bulan awal."""
    n = max(1, min(6, int(n or 1)))
    total = max(total or 0, 0)
    if float(total).is_integer():
        dasar, sisa = divmod(int(total), n)
        nilai = [float(dasar + (1 if i < sisa else 0)) for i in range(n)]
    else:
        dasar = math.floor(total / n * 1000) / 1000
        nilai = [dasar] * n
        nilai[0] = round(total - dasar * (n - 1), 3)
    return nilai + [None] * (6 - n)


def bulan_item(it):
    """Kembalikan enam nilai pemakaian bulanan, otomatis dibagi bila diperlukan."""
    if it["bulan_auto"]:
        return bagi_rata((it["qty"] or 0) + max(it["stok"] or 0, 0), it["bulan_n"])
    return [it[f"b{i}"] for i in range(1, 7)]


def buat_excel(pr, items):
    """Bangun workbook Pendukung PR yang siap diunggah atau diteruskan ke SAP."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Pendukung PR"
    lebar = {"A": 2.1, "B": 13.4, "C": 8, "D": 12.7, "E": 15.4, "F": 43.7, "G": 8.3, "H": 7.4, "I": 18, "J": 9,
             "K": 21.9, "L": 15.7, "M": 11.6, "N": 9.6, "O": 9.6, "P": 9.6, "Q": 9.6, "R": 9.6, "S": 2, "T": 10}
    for k, v in lebar.items():
        ws.column_dimensions[k].width = v
    ws["B2"] = "Pendukung PR (Stock, Terakhir GRN, Periode Pemakaian dan Estimasi GI)"
    ws["B2"].font = Font(bold=True, size=12)
    kepala = [("B", "Date"), ("C", "Plant"), ("D", "PR"), ("E", "Kode Material"), ("F", "Material Description"),
              ("G", "Qty PR"), ("H", "UoM"), ("I", "Estimated Price (IDR)"), ("J", "Stock on Hand"),
              ("K", "Terakhir Purchase\n(Diisi tanggal terakhir GRN)"),
              ("L", "Periode Pemakaian\n(Misal quartal atau semester)")]
    for col, teks in kepala:
        ws.merge_cells(f"{col}3:{col}4")
        ws[f"{col}3"] = teks
    ws.merge_cells("M3:R3")
    ws["M3"] = ("Periode Pengeluaran\n(Diisi estimasi quantity GI dilakukan total Stock on Hand + rencana PR )")
    for i, col in enumerate("MNOPQR", 1):
        ws[f"{col}4"] = f"Bulan ke-{i}"
    ws["T4"] = "crosscek"
    for r in (3, 4):
        ws.row_dimensions[r].height = 32
        for col in "BCDEFGHIJKLMNOPQR":
            c = ws[f"{col}{r}"]
            c.font, c.alignment, c.border, c.fill = Font(bold=True, size=10), TENGAH, KOTAK, KEPALA
    ws["T4"].font = Font(bold=True, size=10)

    tgl = dt.date.fromisoformat(pr["tanggal"]).strftime("%d.%m.%Y")
    baris = 5
    for it in items:
        no = it.get("nomor_sap") or pr["nomor_sap"] or ""
        nomor_sap = int(no) if no.isdigit() else (no or None)
        bulan = bulan_item(it)
        nilai = {"B": tgl, "C": pr["plant"], "D": nomor_sap, "E": it["material"], "F": it["deskripsi"],
                 "G": it["qty"], "H": it["satuan"],
                 "I": f"=G{baris}*{it['harga']:g}" if it["harga"] else None, "J": max(it["stok"] or 0, 0),
                 "K": it["terakhir_purchase"], "L": it["periode"]}
        for i, col in enumerate("MNOPQR"):
            nilai[col] = bulan[i]
        for col, v in nilai.items():
            c = ws[f"{col}{baris}"]
            c.value = v
            c.border = KOTAK
            c.font = Font(size=10)
            if col in "GJMNOPQR":
                c.number_format = "#,##0" if isinstance(v, (int, float)) and float(v).is_integer() else "#,##0.###"
            elif col == "I":
                c.number_format = "#,##0"
        ws[f"T{baris}"] = f"=(G{baris}+J{baris})-SUM(M{baris}:R{baris})"
        baris += 1
    ws[f"F{baris}"] = "SUBTOTAL "
    ws[f"F{baris}"].font = Font(bold=True)
    ws[f"I{baris}"] = f"=SUM(I5:I{baris - 1})" if items else 0
    ws[f"I{baris}"].font = Font(bold=True)
    ws[f"I{baris}"].number_format = "#,##0"
    ws[f"I{baris}"].border = KOTAK
    if pr["judul"] or pr["keterangan"]:
        ws[f"B{baris + 2}"] = " · ".join(x for x in (pr["judul"], pr["keterangan"]) if x)
        ws[f"B{baris + 2}"].font = Font(italic=True, size=9)
    ws.freeze_panes = "E5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"B2:T{baris}"
    ws.print_title_rows = "3:4"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def nama_file(pr, kelompok=None):
    """Susun nama file PR yang aman dan mudah dikenali operator."""
    aman = re.sub(r"[^A-Za-z0-9_-]+", "_", kelompok or pr["judul"] or "")[:30].strip("_")
    return f"Pendukung_PR_{pr['nomor']}{'_' + aman if aman else ''}.xlsx"


def kelompok_default(kategori, material_group, kode=None):
    """Kelompok PR SAP: kategori barang bila diatur, selain itu Material Group (3 digit awal kode)."""
    if kategori:
        return kategori
    grup = material_group or ((kode or "")[:3] if (kode or "")[:3].isdigit() else "")
    return f"Group {grup}" if grup else "Lainnya"
