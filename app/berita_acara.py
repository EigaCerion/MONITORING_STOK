"""Excel "Berita Acara Perhitungan Fisik Persediaan" mengikuti format laporan opname yang sudah dipakai
(sheet per jenis barang, kartu stok vs fisik, selisih lebih/kurang, tanda tangan 3 pihak)."""
import datetime as dt
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TIPIS = Side(style="thin", color="000000")
KOTAK = Border(left=TIPIS, right=TIPIS, top=TIPIS, bottom=TIPIS)
TENGAH = Alignment(horizontal="center", vertical="center", wrap_text=True)
KIRI = Alignment(horizontal="left", vertical="center", wrap_text=True)
KEPALA = PatternFill("solid", fgColor="D9E1F2")
ANGKA = "#,##0.###;-#,##0.###;\"-\""
ANGKA_BULAT = "#,##0;-#,##0;\"-\""


def fmt_angka(v):
    """Format angka qty berita acara tanpa desimal yang tidak diperlukan."""
    return ANGKA_BULAT if isinstance(v, (int, float)) and float(v).is_integer() else ANGKA
LEBAR = {"A": 1.3, "B": 5.3, "C": 14.5, "D": 45, "E": 6, "F": 9, "G": 12.5, "H": 13.7, "I": 8, "J": 8,
         "K": 11, "L": 11, "M": 11, "N": 34}
BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober",
         "November", "Desember"]


def tgl_panjang(iso):
    """Format tanggal ISO menjadi tanggal panjang Indonesia."""
    d = dt.date.fromisoformat(iso)
    return f"{d.day} {BULAN[d.month - 1]} {d.year}"


def selisih(item):
    """Konvensi format lama: Lebih = kartu stok melebihi fisik, Kurang = kartu stok kurang dari fisik."""
    d = round(item["stok_buku"] - item["fisik"], 6)
    return (d, 0.0) if d > 0 else (0.0, -d)


def keterangan(item):
    """Susun keterangan selisih fisik untuk satu item opname."""
    teks = [item["keterangan"]] if item["keterangan"] else []
    lebih, kurang = selisih(item)
    if lebih or kurang:
        if item["trx_sap_posted"]:
            teks.append("Sudah disesuaikan di SAP" + (f" ({item['trx_sap_doc']})" if item["trx_sap_doc"] else ""))
        else:
            teks.append("Belum disesuaikan di SAP")
    return " · ".join(teks)


def _nama_sheet(nama, dipakai):
    """Buat nama worksheet unik yang memenuhi batasan nama Excel."""
    dasar = re.sub(r"[\[\]:*?/\\]", " ", nama).strip()[:31] or "OPNAME"
    hasil, i = dasar, 2
    while hasil.upper() in dipakai:
        hasil = f"{dasar[:28]} {i}"
        i += 1
    dipakai.add(hasil.upper())
    return hasil


def _sel(ws, ref, nilai=None, bold=False, align=None, border=False, fill=None, size=None, fmt=None):
    """Isi dan format satu sel berita acara melalui antarmuka ringkas."""
    c = ws[ref]
    if nilai is not None:
        c.value = nilai
    c.font = Font(bold=bold, size=size or 10)
    if align:
        c.alignment = align
    if border:
        c.border = KOTAK
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt_angka(c.value) if fmt == ANGKA else fmt
    return c


def _sheet_berita_acara(ws, tanggal, jenis, items, st):
    """Bangun sheet berita acara untuk satu jenis barang dan tanggal opname."""
    for col, w in LEBAR.items():
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False

    _sel(ws, "J5", "Nama Perusahaan")
    _sel(ws, "L5", f": {st['nama_perusahaan']}")
    ws.merge_cells("L5:N5")
    ws.merge_cells("B6:I7")
    _sel(ws, "B6", "BERITA ACARA PERHITUNGAN FISIK PERSEDIAAN", bold=True, size=14, align=TENGAH)
    _sel(ws, "J6", "Unit Operasional")
    _sel(ws, "L6", f": {st['unit_operasional']}")
    ws.merge_cells("L6:N6")
    _sel(ws, "J7", "Tanggal Pelaksanaan")
    _sel(ws, "L7", f": {tgl_panjang(tanggal)}")
    ws.merge_cells("L7:N7")
    _sel(ws, "B8", "Jenis Barang       :", bold=True)
    _sel(ws, "D8", jenis, bold=True)

    kepala = [("B10:B11", "No."), ("C10:C11", "Kode Barang"), ("D10:D11", "Nama & Spesifikasi Barang"),
              ("E10:E11", "Uom"), ("F10:F11", "Lokasi"), ("H10:H11", "Stock Opname (Fisik)"),
              ("K10:K11", "Stock yang Benar"), ("N10:N11", "Keterangan")]
    for rng, teks in kepala:
        ws.merge_cells(rng)
        _sel(ws, rng.split(":")[0], teks, bold=True, align=TENGAH)
    _sel(ws, "G10", "Lap.Material", bold=True, align=TENGAH)
    _sel(ws, "G11", "(Kartu Stock)", bold=True, align=TENGAH)
    ws.merge_cells("I10:J10")
    _sel(ws, "I10", "Transaksi Intern", bold=True, align=TENGAH)
    _sel(ws, "I11", "Input", bold=True, align=TENGAH)
    _sel(ws, "J11", "Output", bold=True, align=TENGAH)
    ws.merge_cells("L10:M10")
    _sel(ws, "L10", "Selisih", bold=True, align=TENGAH)
    _sel(ws, "L11", "Lebih", bold=True, align=TENGAH)
    _sel(ws, "M11", "Kurang", bold=True, align=TENGAH)
    for r in (10, 11):
        ws.row_dimensions[r].height = 22
        for col in "BCDEFGHIJKLMN":
            c = ws[f"{col}{r}"]
            c.border = KOTAK
            c.fill = KEPALA

    baris = 12
    for no, it in enumerate(items, 1):
        lebih, kurang = selisih(it)
        nilai = [("B", no, TENGAH, None), ("C", it["material"], KIRI, None), ("D", it["deskripsi"], KIRI, None),
                 ("E", it["satuan"], TENGAH, None), ("F", it["lokasi"] or "", TENGAH, None),
                 ("G", it["stok_buku"], None, ANGKA), ("H", it["fisik"], None, ANGKA), ("I", None, None, ANGKA),
                 ("J", None, None, ANGKA), ("K", it["fisik"], None, ANGKA), ("L", lebih, None, ANGKA),
                 ("M", kurang, None, ANGKA), ("N", keterangan(it), KIRI, None)]
        for col, v, al, fmt in nilai:
            _sel(ws, f"{col}{baris}", v, align=al, border=True, fmt=fmt)
        baris += 1

    ttd = baris + 2
    for (rng, judul), (k_nama, k_jab) in zip(
            [("G{}:H{}", "Dibuat Oleh,"), ("I{}:K{}", "Diperiksa oleh,"), ("L{}:N{}", "Disetujui oleh")],
            [("dibuat_nama", "dibuat_jabatan"), ("diperiksa_nama", "diperiksa_jabatan"),
             ("disetujui_nama", "disetujui_jabatan")]):
        for r, teks, tebal in ((ttd, judul, False), (ttd + 5, st[k_nama], True), (ttd + 6, st[k_jab], False)):
            ref = rng.format(r, r)
            ws.merge_cells(ref)
            c = _sel(ws, ref.split(":")[0], teks, bold=tebal, align=TENGAH)
            if tebal:
                c.font = Font(bold=True, underline="single", size=10)

    ws.freeze_panes = "A12"
    ws.print_area = f"A1:N{ttd + 7}"
    ws.print_title_rows = "10:11"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def _sheet_selisih(ws, items):
    """Bangun sheet ringkasan seluruh item yang memiliki selisih opname."""
    ws.sheet_view.showGridLines = False
    _sel(ws, "B2", "DAFTAR SELISIH STOCK OPNAME", bold=True, size=12)
    kepala = ["No", "Tanggal", "Jenis Barang", "Kode", "Nama Barang", "Uom", "Kartu Stock", "Fisik", "Lebih",
              "Kurang", "Status SAP", "Keterangan"]
    for j, h in enumerate(kepala, 2):
        c = ws.cell(4, j, h)
        c.font, c.alignment, c.border, c.fill = Font(bold=True, size=10), TENGAH, KOTAK, KEPALA
    r = 5
    for no, it in enumerate([i for i in items if any(selisih(i))], 1):
        lebih, kurang = selisih(it)
        nilai = [no, dt.date.fromisoformat(it["tanggal"]), it["jenis_barang"], it["material"], it["deskripsi"],
                 it["satuan"], it["stok_buku"], it["fisik"], lebih, kurang,
                 "Sudah" if it["trx_sap_posted"] else "Belum", it["keterangan"] or ""]
        for j, v in enumerate(nilai, 2):
            c = ws.cell(r, j, v)
            c.border, c.font = KOTAK, Font(size=10)
            if j == 3:
                c.number_format = "dd/mm/yyyy"
            elif 8 <= j <= 11:
                c.number_format = fmt_angka(v)
        r += 1
    if r == 5:
        ws.cell(5, 2, "Tidak ada selisih — semua fisik sama dengan kartu stok.")
    for j, w in enumerate([5, 11, 16, 14, 42, 6, 11, 11, 10, 10, 11, 30], 2):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"


def buat_excel(sesi_items, st):
    """sesi_items: list item (dict) lengkap dengan tanggal, jenis_barang, lokasi, material, deskripsi, satuan, dst."""
    wb = Workbook()
    wb.remove(wb.active)
    grup = {}
    for it in sesi_items:
        grup.setdefault((it["tanggal"], it["jenis_barang"]), []).append(it)
    dipakai = set()
    banyak_tanggal = len({k[0] for k in grup}) > 1
    for (tanggal, jenis), items in sorted(grup.items()):
        nama = f"{jenis} {tanggal[8:10]}-{tanggal[5:7]}" if banyak_tanggal else jenis
        _sheet_berita_acara(wb.create_sheet(_nama_sheet(nama, dipakai)), tanggal, jenis, items, st)
    _sheet_selisih(wb.create_sheet(_nama_sheet("Selisih", dipakai)), sesi_items)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
